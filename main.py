import json
import re
import os
import hashlib
import shutil
import subprocess
import sys
import tempfile
import webbrowser
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from PySide6.QtCore import QPoint, QRect, QSize, Qt, QThread, Signal, QTimer, QUrl
from PySide6.QtGui import QColor, QDragEnterEvent, QDropEvent, QTextCharFormat, QTextCursor, QKeySequence, QShortcut
from PySide6.QtWidgets import (
    QApplication,
    QInputDialog,
    QCheckBox,
    QComboBox,
    QDialog,
    QFileDialog,
    QFrame,
    QGridLayout,
    QHBoxLayout,
    QLabel,
    QLayout,
    QLayoutItem,
    QListWidget,
    QListWidgetItem,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QProgressBar,
    QLineEdit,
    QSizePolicy,
    QSplitter,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

try:
    from PySide6.QtMultimedia import QAudioOutput, QMediaPlayer
except Exception:  # QtMultimedia puede faltar en instalaciones incompletas.
    QAudioOutput = None
    QMediaPlayer = None

from theme import build_stylesheet, load_theme_name, save_theme_name, theme_toggle_text

APP_NAME = "AR_TRANSCRIPTOR"
APP_VERSION = "1.23 importacion estable"
SUPPORTED_EXTENSIONS = {
    ".mp4", ".avi", ".mkv", ".mov", ".webm",
    ".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg", ".wma"
}

LANGUAGES = {
    "Automatico": None,
    "Espanol": "es",
    "Ingles": "en",
    "Portugues": "pt",
    "Italiano": "it",
    "Frances": "fr",
    "Aleman": "de",
    "Otro/manual": "manual",
}

MODELS = ["tiny", "base", "small", "medium", "large", "turbo"]
TASKS = {
    "Transcribir": "transcribe",
    "Traducir al ingles": "translate",
}

FORMAT_MODES = {
    "Texto normal": "normal",
    "P/R alternado simple": "qr_auto",
    "P/R inteligente llamada": "qr_call",
}

# Límites visuales para que una importación grande no congele la interfaz.
MAX_LIST_ITEMS_RENDER = 1500
MAX_SEARCH_HIGHLIGHTS = 800
MAX_IMPORTED_DISPLAY_CHARS = 1_200_000

PERFORMANCE_PROFILES = {
    "Rapido": "fast",
    "Equilibrado": "balanced",
    "Preciso": "accurate",
}

DEVICE_OPTIONS = {
    "Automatico": "auto",
    "CPU": "cpu",
    "GPU CUDA": "cuda",
}

ENGINE_OPTIONS = {
    "Faster Whisper (mas rapido)": "faster",
    "OpenAI Whisper oficial": "openai",
}

COMPUTE_OPTIONS = {
    "Auto": "auto",
    "int8 CPU rapido": "int8",
    "float16 GPU": "float16",
    "float32 compatible": "float32",
}

THREAD_OPTIONS = ["Auto"] + [str(i) for i in range(1, 17)]

CHUNK_OPTIONS = {
    "Sin trocear": 0,
    "5 min": 300,
    "10 min": 600,
    "15 min": 900,
    "30 min": 1800,
}

MODEL_CACHE: Dict[Tuple[str, str, str, str], Any] = {}
AUTOSAVE_SCHEMA_VERSION = 1


def app_data_dir() -> Path:
    """Carpeta local de la app para copias discretas de transcripciones."""
    root = os.environ.get("LOCALAPPDATA") or os.environ.get("APPDATA")
    if root:
        return Path(root) / APP_NAME
    return Path.home() / f".{APP_NAME.lower()}"


def autosave_dir() -> Path:
    return app_data_dir() / "copias_transcripciones"


def safe_file_slug(value: str, fallback: str = "transcripcion") -> str:
    value = re.sub(r"[^A-Za-z0-9_\-]+", "_", value.strip())
    value = value.strip("._-")
    return (value or fallback)[:70]


def default_settings() -> Dict[str, Any]:
    return {
        "model": "small",
        "language": "Espanol",
        "manual_language": "",
        "task": "Transcribir",
        "format": "Texto normal",
        "performance": "Rapido",
        "device": "Automatico",
        "engine": "Faster Whisper (mas rapido)",
        "compute": "Auto",
        "optimize_audio": True,
        "clean_text": True,
        "threads": "Auto",
        "chunk": "10 min",
    }


def role_from_label(label: str) -> str:
    value = _remove_accents_for_match(str(label or "")).strip().lower()
    if value in {"pregunta", "operador", "operadora", "entrevistador", "entrevistadora", "personal policial", "policia"}:
        return "question"
    if value in {"respuesta", "llamante", "requirente", "entrevistado", "entrevistada", "denunciante"}:
        return "answer"
    return "question" if "preg" in value or "oper" in value else "answer"


def seconds_to_hhmmss(seconds: float) -> str:
    seconds = max(float(seconds or 0), 0.0)
    total = int(seconds)
    h = total // 3600
    m = (total % 3600) // 60
    s = total % 60
    return f"{h:02d}:{m:02d}:{s:02d}"


def seconds_to_duration_label(seconds: Optional[float]) -> str:
    """Duración compacta para mostrar en pantalla."""
    if seconds is None:
        return "--:--"
    try:
        seconds = float(seconds)
    except (TypeError, ValueError):
        return "--:--"
    if seconds <= 0:
        return "--:--"
    total = int(round(seconds))
    h = total // 3600
    m = (total % 3600) // 60
    s = total % 60
    if h:
        return f"{h:02d}:{m:02d}:{s:02d}"
    return f"{m:02d}:{s:02d}"


def probe_media_duration_seconds(file_path: str) -> Optional[float]:
    """Obtiene duración del audio/video con ffprobe sin subir nada a internet."""
    if not file_path:
        return None
    ffprobe = shutil.which("ffprobe")
    if not ffprobe:
        return None
    try:
        completed = subprocess.run(
            [
                ffprobe,
                "-v", "error",
                "-show_entries", "format=duration",
                "-of", "default=noprint_wrappers=1:nokey=1",
                file_path,
            ],
            capture_output=True,
            text=True,
            timeout=8,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
        )
        if completed.returncode != 0:
            return None
        lines = (completed.stdout or "").strip().splitlines()
        if not lines:
            return None
        duration = float(lines[0].strip())
        return duration if duration > 0 else None
    except Exception:
        return None


def duration_from_segments(segments: List[Dict[str, Any]]) -> Optional[float]:
    ends = []
    for seg in segments or []:
        try:
            ends.append(float(seg.get("end", 0) or 0))
        except (TypeError, ValueError):
            pass
    return max(ends) if ends else None


def normalize_word_items(words: Any, offset: float = 0.0) -> List[Dict[str, Any]]:
    """Normaliza palabras con timestamp provenientes de Faster-Whisper u OpenAI Whisper."""
    normalized: List[Dict[str, Any]] = []
    for word in words or []:
        try:
            if isinstance(word, dict):
                text = str(word.get("word", word.get("text", ""))).strip()
                start = float(word.get("start", 0) or 0) + offset
                end = float(word.get("end", start) or start) + offset
                probability = word.get("probability", word.get("prob", None))
            else:
                text = str(getattr(word, "word", getattr(word, "text", ""))).strip()
                start = float(getattr(word, "start", 0) or 0) + offset
                end = float(getattr(word, "end", start) or start) + offset
                probability = getattr(word, "probability", getattr(word, "prob", None))
            if not text:
                continue
            item: Dict[str, Any] = {"word": text, "start": start, "end": end}
            if probability is not None:
                try:
                    item["probability"] = float(probability)
                except (TypeError, ValueError):
                    pass
            normalized.append(item)
        except Exception:
            continue
    return normalized


def words_for_time(words: List[Dict[str, Any]], pos_s: float) -> Optional[int]:
    for idx, word in enumerate(words or []):
        try:
            start = float(word.get("start", 0) or 0)
            end = float(word.get("end", start) or start)
        except (TypeError, ValueError):
            continue
        if start <= pos_s <= max(end, start + 0.08):
            return idx
    return None


def seconds_to_srt_time(seconds: float) -> str:
    seconds = max(float(seconds or 0), 0.0)
    ms_total = int(round(seconds * 1000))
    h = ms_total // 3_600_000
    ms_total %= 3_600_000
    m = ms_total // 60_000
    ms_total %= 60_000
    s = ms_total // 1000
    ms = ms_total % 1000
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def seconds_to_vtt_time(seconds: float) -> str:
    return seconds_to_srt_time(seconds).replace(",", ".")


def make_plain_segments(segments: List[Dict[str, Any]]) -> str:
    lines = []
    for seg in segments:
        start = seconds_to_hhmmss(seg.get("start", 0))
        end = seconds_to_hhmmss(seg.get("end", 0))
        text = str(seg.get("text", "")).strip()
        lines.append(f"[{start} - {end}] {text}")
    return "\n".join(lines)


def _remove_accents_for_match(text: str) -> str:
    replacements = {
        "á": "a", "é": "e", "í": "i", "ó": "o", "ú": "u", "ü": "u", "ñ": "n",
        "Á": "A", "É": "E", "Í": "I", "Ó": "O", "Ú": "U", "Ü": "U", "Ñ": "N",
    }
    for src, dst in replacements.items():
        text = text.replace(src, dst)
    return text


def basic_police_cleanup(text: str) -> str:
    """
    Limpieza simple para audios policiales/llamadas.
    No reemplaza la escucha humana: solo corrige errores repetidos comunes.
    """
    if not text:
        return ""
    cleaned = text.strip()
    replacements = [
        (r"\b[Pp]olic[ií]a\s+de\s+Tral[oó]n\b", "Policía de Trelew"),
        (r"\b[Tt]ral[oó]n\b", "Trelew"),
        (r"[¿?]?\s*[Ss][oó]pezido\s*[¿?]?", "¿Su apellido?"),
        (r"[¿?]?\s*[Ss]opesido\s*[¿?]?", "¿Su apellido?"),
        (r"[¿?]?\s*[Ss]upezido\s*[¿?]?", "¿Su apellido?"),
        (r"\b[Ss]ópezido\b", "su apellido"),
        (r"\b[Ss]opesido\b", "su apellido"),
        (r"\b[Mm]oreira,\s*¿Su apellido\?", "Moreira. ¿Su apellido?"),
    ]
    for pattern, replacement in replacements:
        cleaned = re.sub(pattern, replacement, cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    cleaned = cleaned.replace(" ¿", " ¿").replace("? ?", "?")
    return cleaned


def _repetition_key(text: str) -> str:
    key = _remove_accents_for_match(text).lower()
    key = re.sub(r"[^a-z0-9]+", " ", key).strip()
    return key


def _prepare_segments_for_pr(segments: List[Dict[str, Any]], cleanup: bool = True) -> Tuple[List[Dict[str, Any]], List[str]]:
    prepared: List[Dict[str, Any]] = []
    notes: List[str] = []
    last_key = ""
    repeated_count = 0
    omitted_repetitions = 0

    for seg in segments:
        text = str(seg.get("text", "")).strip()
        if cleanup:
            text = basic_police_cleanup(text)
        key = _repetition_key(text)
        if key and key == last_key and len(key) <= 35:
            repeated_count += 1
        else:
            repeated_count = 1
            last_key = key

        # Whisper a veces hallucina una frase corta repetida al final del audio.
        # Conservamos hasta 2 apariciones y marcamos el resto como observacion.
        if key and repeated_count > 2 and len(text.split()) <= 4:
            omitted_repetitions += 1
            continue

        copied = dict(seg)
        copied["text"] = text
        prepared.append(copied)

    if omitted_repetitions:
        notes.append(f"Se omitieron {omitted_repetitions} segmento/s repetido/s probablemente generados por ruido, silencio o audio no inteligible. Revisar escucha original.")
    return prepared, notes


def format_pr_blocks(blocks: List[Dict[str, Any]], notes: Optional[List[str]] = None, include_auto_note: bool = False) -> str:
    lines: List[str] = []
    for block in blocks:
        label = str(block.get("label", "PREGUNTA")).strip().upper() or "PREGUNTA"
        start = seconds_to_hhmmss(block.get("start", 0))
        end = seconds_to_hhmmss(block.get("end", 0))
        text = str(block.get("text", "")).strip()
        if start == "00:00:00" and end == "00:00:00":
            lines.append(f"{label}: {text}".rstrip())
        else:
            lines.append(f"{label}: [{start} - {end}] {text}".rstrip())
    if notes:
        lines.append("OBSERVACIONES DE TRANSCRIPCION: " + " ".join(notes))
    if include_auto_note:
        lines.append("NOTA: Formato generado automaticamente. Verificar voces, nombres propios, domicilios y frases dudosas contra el audio original.")
    return "\n\n".join(lines).strip()


def build_qr_blocks(segments: List[Dict[str, Any]], fallback_text: str = "", cleanup: bool = True) -> Tuple[List[Dict[str, Any]], List[str]]:
    if not segments:
        text = basic_police_cleanup(fallback_text.strip()) if cleanup else fallback_text.strip()
        if not text:
            return [
                {"role": "question", "label": "PREGUNTA", "start": 0, "end": 0, "text": ""},
                {"role": "answer", "label": "RESPUESTA", "start": 0, "end": 0, "text": ""},
            ], []
        return [{"role": "question", "label": "PREGUNTA", "start": 0, "end": 0, "text": text}], []

    prepared, notes = _prepare_segments_for_pr(segments, cleanup=cleanup)
    blocks: List[Dict[str, Any]] = []
    for idx, seg in enumerate(prepared):
        role = "question" if idx % 2 == 0 else "answer"
        blocks.append({
            "role": role,
            "label": "PREGUNTA" if role == "question" else "RESPUESTA",
            "start": seg.get("start", 0),
            "end": seg.get("end", 0),
            "text": str(seg.get("text", "")).strip(),
            "words": normalize_word_items(seg.get("words", [])),
        })
    return blocks, notes


def build_qr_text(segments: List[Dict[str, Any]], fallback_text: str = "", cleanup: bool = True) -> str:
    """
    Genera un borrador editable en formato PREGUNTA/RESPUESTA alternando segmentos.
    Es util solo como base rapida; para llamadas conviene P/R inteligente llamada.
    """
    blocks, notes = build_qr_blocks(segments, fallback_text, cleanup=cleanup)
    return format_pr_blocks(blocks, notes)


def _score_operator(text: str) -> int:
    t = _remove_accents_for_match(text).lower()
    score = 0
    operator_phrases = [
        "policia", "si senora", "senora", "donde seria", "direccion", "que barrio",
        "apellido", "su apellido", "listo", "mandamos", "movil", "ahi esta yendo",
        "a donde", "lote 11 manzana", "manzana?", "como?", "si?",
    ]
    for phrase in operator_phrases:
        if phrase in t:
            score += 2
    if "¿" in text or "?" in text:
        score += 1
    return score


def _score_requirente(text: str) -> int:
    t = _remove_accents_for_match(text).lower()
    score = 0
    requirente_phrases = [
        "me robaron", "robado", "chorros", "me tiraron", "tiro", "mis nenes",
        "estoy con", "adentro de la casa", "por favor", "paso de indios",
        "lote", "manzana", "moreira", "yo te llame", "hola", "que tal",
    ]
    for phrase in requirente_phrases:
        if phrase in t:
            score += 2
    # Respuestas cortas tipicas de datos: direccion, barrio, nombre, numero.
    if re.search(r"\b\d{1,4}\b", t):
        score += 1
    return score


def _classify_call_segment(text: str, previous: Optional[str], last_operator_question: bool) -> str:
    clean = text.strip()
    if not clean:
        return previous or "RESPUESTA"
    op = _score_operator(clean)
    req = _score_requirente(clean)
    words = clean.split()
    t = _remove_accents_for_match(clean).lower()

    if t.startswith("policia"):
        return "PREGUNTA"
    if "listo" in t or "mandamos" in t or "movil" in t:
        return "PREGUNTA"
    if "yo te llame" in t or "me tiraron" in t or "mis nenes" in t:
        return "RESPUESTA"

    # Si el operador acaba de preguntar y viene una respuesta corta de datos, asignarla al requirente.
    if last_operator_question and len(words) <= 8 and req >= op:
        return "RESPUESTA"

    # Preguntas muy cortas tipo '¿Sí?', '¿Cómo?', '¿A dónde señora?' suelen ser del operador.
    if ("¿" in clean or "?" in clean) and op >= req:
        return "PREGUNTA"

    if op > req:
        return "PREGUNTA"
    if req > op:
        return "RESPUESTA"

    # Alternancia prudente si no hay suficientes pistas.
    if previous == "PREGUNTA" and len(words) <= 10:
        return "RESPUESTA"
    if previous == "RESPUESTA" and ("¿" in clean or "?" in clean):
        return "PREGUNTA"
    return previous or "RESPUESTA"


def build_call_pr_blocks(segments: List[Dict[str, Any]], fallback_text: str = "", cleanup: bool = True) -> Tuple[List[Dict[str, Any]], List[str]]:
    """
    Crea bloques editables PREGUNTA/RESPUESTA para llamadas/entrevistas breves.
    No es diarizacion real: siempre debe revisarse escuchando el audio.
    """
    if not segments:
        text = basic_police_cleanup(fallback_text.strip()) if cleanup else fallback_text.strip()
        if not text:
            return [
                {"role": "question", "label": "PREGUNTA", "start": 0, "end": 0, "text": ""},
                {"role": "answer", "label": "RESPUESTA", "start": 0, "end": 0, "text": ""},
            ], []
        return [{"role": "answer", "label": "RESPUESTA", "start": 0, "end": 0, "text": text}], []

    prepared, notes = _prepare_segments_for_pr(segments, cleanup=cleanup)
    blocks: List[Dict[str, Any]] = []
    previous_label: Optional[str] = None
    last_operator_question = False

    for seg in prepared:
        text = str(seg.get("text", "")).strip()
        if not text:
            continue
        label = _classify_call_segment(text, previous_label, last_operator_question)
        last_operator_question = label == "PREGUNTA" and ("¿" in text or "?" in text or _score_operator(text) > 0)

        seg_words = normalize_word_items(seg.get("words", []))
        if blocks and blocks[-1]["label"] == label:
            blocks[-1]["end"] = seg.get("end", blocks[-1]["end"])
            blocks[-1]["text"] = re.sub(r"\s+", " ", (blocks[-1]["text"] + " " + text)).strip()
            blocks[-1].setdefault("words", [])
            blocks[-1]["words"].extend(seg_words)
        else:
            blocks.append({
                "role": "question" if label == "PREGUNTA" else "answer",
                "label": label,
                "start": seg.get("start", 0),
                "end": seg.get("end", 0),
                "text": text,
                "words": seg_words,
            })
        previous_label = label

    return blocks, notes


def build_call_pr_text(segments: List[Dict[str, Any]], fallback_text: str = "", cleanup: bool = True) -> str:
    blocks, notes = build_call_pr_blocks(segments, fallback_text, cleanup=cleanup)
    return format_pr_blocks(blocks, notes, include_auto_note=True)


def build_formatted_text(mode: str, segments: List[Dict[str, Any]], fallback_text: str = "", cleanup: bool = True) -> str:
    if mode == "qr_auto":
        return build_qr_text(segments, fallback_text, cleanup=cleanup)
    if mode == "qr_call":
        return build_call_pr_text(segments, fallback_text, cleanup=cleanup)
    return basic_police_cleanup(fallback_text) if cleanup else fallback_text


def build_txt(metadata: Dict[str, Any], text: str, segments: List[Dict[str, Any]]) -> str:
    header = [
        "TRANSCRIPCION DE AUDIO/VIDEO",
        "",
        f"Archivo: {metadata.get('archivo', '')}",
        f"Modelo: {metadata.get('modelo', '')}",
        f"Idioma: {metadata.get('idioma', '')}",
        f"Tarea: {metadata.get('tarea', '')}",
        f"Formato: {metadata.get('formato', '')}",
        f"Perfil rendimiento: {metadata.get('perfil_rendimiento', '')}",
        f"Dispositivo: {metadata.get('dispositivo', '')}",
        f"Motor: {metadata.get('motor', '')}",
        f"Precision/compute: {metadata.get('precision_compute', '')}",
        f"Audio optimizado: {metadata.get('audio_optimizado', '')}",
        f"Hilos CPU/FFmpeg: {metadata.get('hilos_cpu_ffmpeg', '')}",
        f"Troceo archivo: {metadata.get('troceo_archivo', '')}",
        f"Limpieza basica: {metadata.get('limpieza_basica', '')}",
        f"Fecha/hora: {metadata.get('fecha_hora', '')}",
        "",
        "TEXTO COMPLETO:",
        text.strip(),
        "",
        "SEGMENTOS:",
        make_plain_segments(segments),
        "",
    ]
    return "\n".join(header)


def build_srt(segments: List[Dict[str, Any]]) -> str:
    blocks = []
    for idx, seg in enumerate(segments, start=1):
        start = seconds_to_srt_time(seg.get("start", 0))
        end = seconds_to_srt_time(seg.get("end", 0))
        text = str(seg.get("text", "")).strip()
        blocks.append(f"{idx}\n{start} --> {end}\n{text}")
    return "\n\n".join(blocks) + ("\n" if blocks else "")


def build_vtt(segments: List[Dict[str, Any]]) -> str:
    blocks = ["WEBVTT", ""]
    for seg in segments:
        start = seconds_to_vtt_time(seg.get("start", 0))
        end = seconds_to_vtt_time(seg.get("end", 0))
        text = str(seg.get("text", "")).strip()
        blocks.append(f"{start} --> {end}\n{text}\n")
    return "\n".join(blocks)


def safe_cpu_threads() -> int:
    count = os.cpu_count() or 2
    # Deja algo de margen para que la interfaz no quede congelada en equipos chicos.
    return max(1, min(count - 1 if count > 2 else count, 8))


def resolve_thread_count(value: int) -> int:
    if value and value > 0:
        return max(1, min(int(value), 16))
    return safe_cpu_threads()


class TranscriptionWorker(QThread):
    status_changed = Signal(str)
    progress_changed = Signal(int)
    finished_ok = Signal(dict)
    failed = Signal(str)

    def __init__(
        self,
        file_path: str,
        model_name: str,
        language: Optional[str],
        task: str,
        performance_profile: str,
        device_mode: str,
        engine: str,
        compute_type: str,
        optimize_audio: bool,
        cpu_threads: int,
        chunk_seconds: int,
    ):
        super().__init__()
        self.file_path = file_path
        self.model_name = model_name
        self.language = language
        self.task = task
        self.performance_profile = performance_profile
        self.device_mode = device_mode
        self.engine = engine
        self.compute_type = compute_type
        self.optimize_audio = optimize_audio
        self.cpu_threads = resolve_thread_count(cpu_threads)
        self.chunk_seconds = int(chunk_seconds or 0)
        self._temp_dir: Optional[tempfile.TemporaryDirectory] = None

    def _resolve_device(self) -> str:
        if self.device_mode == "cpu":
            return "cpu"
        if self.device_mode == "cuda":
            return "cuda"
        try:
            import torch
            return "cuda" if torch.cuda.is_available() else "cpu"
        except Exception:
            return "cpu"

    def _ensure_temp_dir(self) -> Path:
        if self._temp_dir is None:
            self._temp_dir = tempfile.TemporaryDirectory(prefix="whisper_dpi_")
        return Path(self._temp_dir.name)

    def _ffmpeg_base(self) -> List[str]:
        cmd = ["ffmpeg", "-y", "-hide_banner", "-loglevel", "error"]
        if self.cpu_threads > 0:
            cmd += ["-threads", str(self.cpu_threads)]
        return cmd

    def _prepare_audio_inputs(self) -> List[str]:
        """
        Prepara el archivo para Whisper.
        - Sin troceo: opcionalmente convierte a WAV mono 16 kHz.
        - Con troceo: genera partes WAV mono 16 kHz, util para entrevistas/videos largos.
        """
        ffmpeg_available = shutil.which("ffmpeg") is not None
        if not ffmpeg_available:
            return [self.file_path]

        if self.chunk_seconds > 0:
            temp_root = self._ensure_temp_dir()
            chunks_dir = temp_root / "chunks"
            chunks_dir.mkdir(parents=True, exist_ok=True)
            pattern = str(chunks_dir / "chunk_%05d.wav")
            cmd = self._ffmpeg_base() + [
                "-i", self.file_path,
                "-vn",
                "-ac", "1",
                "-ar", "16000",
                "-f", "segment",
                "-segment_time", str(self.chunk_seconds),
                "-reset_timestamps", "1",
                pattern,
            ]
            self.status_changed.emit(f"Troceando audio/video cada {self.chunk_seconds // 60} min con FFmpeg...")
            self.progress_changed.emit(16)
            completed = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
            if completed.returncode != 0:
                err = completed.stderr.strip() or "FFmpeg no pudo trocear el archivo."
                raise RuntimeError(err)
            chunks = sorted(str(p) for p in chunks_dir.glob("chunk_*.wav"))
            if not chunks:
                raise RuntimeError("FFmpeg no genero fragmentos de audio.")
            return chunks

        if not self.optimize_audio:
            return [self.file_path]

        temp_root = self._ensure_temp_dir()
        output_path = str(temp_root / "audio_16k_mono.wav")
        cmd = self._ffmpeg_base() + [
            "-i", self.file_path,
            "-vn",
            "-ac", "1",
            "-ar", "16000",
            "-f", "wav",
            output_path,
        ]
        self.status_changed.emit("Optimizando audio/video para Whisper...")
        self.progress_changed.emit(18)
        completed = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
        if completed.returncode != 0:
            err = completed.stderr.strip() or "FFmpeg no pudo convertir el archivo."
            raise RuntimeError(err)
        return [output_path]

    def _merge_chunk_result(self, result: Dict[str, Any], offset: float) -> Tuple[str, List[Dict[str, Any]]]:
        text = str(result.get("text", "")).strip()
        merged_segments: List[Dict[str, Any]] = []
        for seg in result.get("segments", []) or []:
            copied = dict(seg)
            copied["start"] = float(copied.get("start", 0) or 0) + offset
            copied["end"] = float(copied.get("end", 0) or 0) + offset
            copied["words"] = normalize_word_items(copied.get("words", []), offset=offset)
            merged_segments.append(copied)
        return text, merged_segments

    def _transcribe_options(self, device: str) -> Dict[str, Any]:
        options: Dict[str, Any] = {
            "task": self.task,
            "verbose": False,
            "fp16": device == "cuda",
            "word_timestamps": True,
        }
        if self.language:
            options["language"] = self.language

        # Perfiles simples para el operador: mas rapido o mas preciso.
        if self.performance_profile == "fast":
            options.update({
                "beam_size": 1,
                "best_of": 1,
                "temperature": 0,
                "condition_on_previous_text": False,
            })
        elif self.performance_profile == "balanced":
            options.update({
                "beam_size": 3,
                "best_of": 3,
                "temperature": 0,
                "condition_on_previous_text": False,
            })
        elif self.performance_profile == "accurate":
            options.update({
                "beam_size": 5,
                "best_of": 5,
                "temperature": 0,
                "condition_on_previous_text": True,
            })
        return options

    def _resolve_compute_type(self, device: str) -> str:
        if self.compute_type and self.compute_type != "auto":
            return self.compute_type
        return "float16" if device == "cuda" else "int8"

    def _faster_transcribe_one(self, model: Any, transcribe_path: str, options: Dict[str, Any]) -> Dict[str, Any]:
        kwargs: Dict[str, Any] = {
            "language": self.language,
            "task": self.task,
            "beam_size": int(options.get("beam_size", 1)),
            "best_of": int(options.get("best_of", 1)),
            "temperature": 0,
            "condition_on_previous_text": bool(options.get("condition_on_previous_text", False)),
            "vad_filter": False,
            "word_timestamps": True,
        }
        if not self.language:
            kwargs.pop("language", None)
        segments_iter, info = model.transcribe(transcribe_path, **kwargs)
        segments: List[Dict[str, Any]] = []
        parts: List[str] = []
        for seg in segments_iter:
            text = str(getattr(seg, "text", "")).strip()
            if text:
                parts.append(text)
            segments.append({
                "start": float(getattr(seg, "start", 0.0) or 0.0),
                "end": float(getattr(seg, "end", 0.0) or 0.0),
                "text": text,
                "words": normalize_word_items(getattr(seg, "words", []) or []),
            })
        return {"text": " ".join(parts).strip(), "segments": segments, "language": getattr(info, "language", None)}

    def run(self) -> None:
        try:
            if self.isInterruptionRequested():
                self.status_changed.emit("Cancelado antes de iniciar.")
                return

            self.status_changed.emit("Preparando motor de transcripcion...")
            self.progress_changed.emit(8)
            # Ajusta hilos antes de que PyTorch/CT2 empiecen a trabajar.
            os.environ["OMP_NUM_THREADS"] = str(self.cpu_threads)
            os.environ["MKL_NUM_THREADS"] = str(self.cpu_threads)
            os.environ["NUMEXPR_NUM_THREADS"] = str(self.cpu_threads)

            try:
                import torch
                if self.device_mode in {"auto", "cpu"}:
                    torch.set_num_threads(self.cpu_threads)
                    try:
                        torch.set_num_interop_threads(max(1, min(2, self.cpu_threads)))
                    except Exception:
                        pass
            except Exception:
                pass

            if self.isInterruptionRequested():
                self.status_changed.emit("Cancelado antes de cargar modelo.")
                return

            device = self._resolve_device()
            compute = self._resolve_compute_type(device)
            cache_key = (self.engine, self.model_name, device, compute)
            if cache_key in MODEL_CACHE:
                self.status_changed.emit(f"Usando modelo '{self.model_name}' ya cargado ({self.engine}, {device.upper()}, {compute})...")
                model = MODEL_CACHE[cache_key]
                self.progress_changed.emit(30)
            else:
                if self.engine == "faster":
                    self.status_changed.emit(f"Cargando Faster Whisper '{self.model_name}' en {device.upper()} ({compute})...")
                    self.progress_changed.emit(25)
                    try:
                        from faster_whisper import WhisperModel
                    except Exception as exc:
                        raise RuntimeError("No esta instalado faster-whisper. Ejecuta install_dependencies.bat de esta version o cambia el motor a OpenAI Whisper oficial.") from exc
                    model = WhisperModel(
                        self.model_name,
                        device=device,
                        compute_type=compute,
                        cpu_threads=self.cpu_threads if device == "cpu" else 0,
                        num_workers=1,
                    )
                else:
                    self.status_changed.emit(f"Cargando OpenAI Whisper '{self.model_name}' en {device.upper()}...")
                    self.progress_changed.emit(25)
                    import whisper  # Import interno para que la app abra aunque falte instalar Whisper.
                    model = whisper.load_model(self.model_name, device=device)
                MODEL_CACHE[cache_key] = model

            if self.isInterruptionRequested():
                self.status_changed.emit("Cancelado antes de transcribir.")
                return

            input_paths = self._prepare_audio_inputs()
            options = self._transcribe_options(device)

            all_text: List[str] = []
            all_segments: List[Dict[str, Any]] = []
            raw_chunks: List[Dict[str, Any]] = []
            total_chunks = len(input_paths)
            for idx, transcribe_path in enumerate(input_paths):
                if self.isInterruptionRequested():
                    self.status_changed.emit("Transcripcion cancelada. Resultado descartado.")
                    return

                if total_chunks > 1:
                    self.status_changed.emit(f"Transcribiendo fragmento {idx + 1}/{total_chunks}...")
                else:
                    self.status_changed.emit("Transcribiendo...")
                start_progress = 40
                end_progress = 98
                self.progress_changed.emit(start_progress + int((idx / max(total_chunks, 1)) * (end_progress - start_progress)))

                if self.engine == "faster":
                    chunk_result = self._faster_transcribe_one(model, transcribe_path, options)
                else:
                    chunk_result = model.transcribe(transcribe_path, **options)
                offset = float(idx * self.chunk_seconds) if self.chunk_seconds > 0 else 0.0
                chunk_text, chunk_segments = self._merge_chunk_result(chunk_result, offset)
                if chunk_text:
                    all_text.append(chunk_text)
                all_segments.extend(chunk_segments)
                raw_chunks.append({
                    "index": idx,
                    "path": Path(transcribe_path).name,
                    "offset_seconds": offset,
                    "text": chunk_text,
                    "segments_count": len(chunk_segments),
                })

            result = {
                "text": " ".join(all_text).strip(),
                "segments": all_segments,
                "language": None,
                "_dpi_chunks": raw_chunks,
                "_dpi_runtime": {
                    "device": device,
                    "profile": self.performance_profile,
                    "engine": self.engine,
                    "compute_type": compute,
                    "optimized_audio": bool(self.optimize_audio or self.chunk_seconds > 0),
                    "model_cached": True,
                    "cpu_threads": self.cpu_threads,
                    "chunk_seconds": self.chunk_seconds,
                    "chunks": total_chunks,
                },
            }

            if self.isInterruptionRequested():
                self.status_changed.emit("Transcripcion cancelada. Resultado descartado.")
                return

            self.progress_changed.emit(100)
            self.status_changed.emit("Finalizado.")
            self.finished_ok.emit(result)
        except Exception as exc:
            self.failed.emit(str(exc))
        finally:
            if self._temp_dir is not None:
                try:
                    self._temp_dir.cleanup()
                except Exception:
                    pass


class FlowLayout(QLayout):
    def __init__(self, parent: Optional[QWidget] = None, margin: int = 0, h_spacing: int = 8, v_spacing: int = 8):
        super().__init__(parent)
        self._items: List[QLayoutItem] = []
        self._h_spacing = h_spacing
        self._v_spacing = v_spacing
        self.setContentsMargins(margin, margin, margin, margin)

    def addItem(self, item: QLayoutItem) -> None:
        self._items.append(item)

    def count(self) -> int:
        return len(self._items)

    def itemAt(self, index: int) -> Optional[QLayoutItem]:
        if 0 <= index < len(self._items):
            return self._items[index]
        return None

    def takeAt(self, index: int) -> Optional[QLayoutItem]:
        if 0 <= index < len(self._items):
            return self._items.pop(index)
        return None

    def expandingDirections(self):
        return Qt.Orientation(0)

    def hasHeightForWidth(self) -> bool:
        return True

    def heightForWidth(self, width: int) -> int:
        return self._do_layout(QRect(0, 0, width, 0), test_only=True)

    def setGeometry(self, rect: QRect) -> None:
        super().setGeometry(rect)
        self._do_layout(rect, test_only=False)

    def sizeHint(self) -> QSize:
        return self.minimumSize()

    def minimumSize(self) -> QSize:
        size = QSize()
        for item in self._items:
            size = size.expandedTo(item.minimumSize())
        margins = self.contentsMargins()
        size += QSize(margins.left() + margins.right(), margins.top() + margins.bottom())
        return size

    def _do_layout(self, rect: QRect, test_only: bool) -> int:
        margins = self.contentsMargins()
        effective = rect.adjusted(margins.left(), margins.top(), -margins.right(), -margins.bottom())
        x = effective.x()
        y = effective.y()
        line_height = 0

        for item in self._items:
            widget = item.widget()
            if widget is not None and not widget.isVisible():
                continue
            next_x = x + item.sizeHint().width() + self._h_spacing
            if next_x - self._h_spacing > effective.right() and line_height > 0:
                x = effective.x()
                y += line_height + self._v_spacing
                next_x = x + item.sizeHint().width() + self._h_spacing
                line_height = 0
            if not test_only:
                item.setGeometry(QRect(QPoint(x, y), item.sizeHint()))
            x = next_x
            line_height = max(line_height, item.sizeHint().height())

        return y + line_height - rect.y() + margins.bottom()


class DropFrame(QFrame):
    file_dropped = Signal(str)

    def __init__(self):
        super().__init__()
        self.setAcceptDrops(True)
        self.setObjectName("dropFrame")
        layout = QVBoxLayout(self)
        layout.setContentsMargins(12, 8, 12, 8)
        layout.setSpacing(3)
        # Zona de carga: drop area ampliada con icono y formatos aceptados.
        icon = QLabel("\U0001F4C1")
        icon.setAlignment(Qt.AlignCenter)
        icon.setObjectName("dropIcon")
        title = QLabel("Arrastra un audio/video aca")
        title.setAlignment(Qt.AlignCenter)
        title.setObjectName("dropTitle")
        subtitle = QLabel("o usa el boton Seleccionar archivo")
        subtitle.setAlignment(Qt.AlignCenter)
        subtitle.setObjectName("dropSubtitle")
        formats = QLabel("MP4, AVI, MKV, MOV, WEBM, MP3, WAV, M4A, AAC, FLAC, OGG, WMA")
        formats.setAlignment(Qt.AlignCenter)
        formats.setWordWrap(True)
        formats.setObjectName("dropFormats")
        layout.addWidget(icon)
        layout.addWidget(title)
        layout.addWidget(subtitle)
        layout.addWidget(formats)

    def dragEnterEvent(self, event: QDragEnterEvent) -> None:
        if event.mimeData().hasUrls():
            for url in event.mimeData().urls():
                path = Path(url.toLocalFile())
                if path.suffix.lower() in SUPPORTED_EXTENSIONS:
                    event.acceptProposedAction()
                    return
        event.ignore()

    def dropEvent(self, event: QDropEvent) -> None:
        for url in event.mimeData().urls():
            path = Path(url.toLocalFile())
            if path.suffix.lower() in SUPPORTED_EXTENSIONS:
                self.file_dropped.emit(str(path))
                event.acceptProposedAction()
                return
        event.ignore()


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"{APP_NAME} - {APP_VERSION}")
        self.resize(1320, 720)
        self.setMinimumSize(980, 600)
        self.theme_name = load_theme_name()
        self.settings = default_settings()
        self.selected_file: Optional[str] = None
        self.last_result: Optional[Dict[str, Any]] = None
        self.pr_blocks: List[Dict[str, Any]] = []
        self.workspace_mode = "pro"
        self.pr_notes: List[str] = []
        self.current_backup_path: Optional[Path] = None
        self.media_duration_seconds: Optional[float] = None
        self.search_matches: List[Tuple[int, int]] = []
        self.search_match_index: int = -1
        self.speaker_names: Dict[str, str] = {"SPEAKER_1": "Speaker 1", "SPEAKER_2": "Speaker 2"}
        self.review_keywords: List[str] = []
        self.review_findings: List[Dict[str, Any]] = []
        self._bulk_loading_ui: bool = False
        self.worker: Optional[TranscriptionWorker] = None
        self._setup_ui()
        self._apply_style()
        self._check_environment()

    def _setup_ui(self) -> None:
        root = QWidget()
        self.setCentralWidget(root)
        main = QVBoxLayout(root)
        main.setContentsMargins(10, 8, 10, 8)
        main.setSpacing(7)

        header = QHBoxLayout()
        title_box = QVBoxLayout()
        title_box.setSpacing(2)
        # Header visual: marca, estado local y switch persistente de tema.
        title = QLabel(f"\U0001F399\ufe0f {APP_NAME}")
        title.setObjectName("title")
        subtitle = QLabel(f"Transcripcion local de audio y video con OpenAI Whisper - {APP_VERSION}")
        subtitle.setObjectName("subtitle")
        title_box.addWidget(title)
        title_box.addWidget(subtitle)
        header.addLayout(title_box)
        header.addStretch(1)
        badge = QLabel("LOCAL / SIN SUBIR ARCHIVOS")
        badge.setObjectName("badge")
        header.addWidget(badge)
        self.theme_toggle_button = QPushButton()
        self.theme_toggle_button.setObjectName("themeToggle")
        self.theme_toggle_button.setToolTip("Cambiar modo claro/oscuro")
        self.theme_toggle_button.clicked.connect(self.toggle_theme)
        header.addWidget(self.theme_toggle_button)
        main.addLayout(header)

        header_separator = QFrame()
        header_separator.setObjectName("headerSeparator")
        main.addWidget(header_separator)

        top_card = QFrame()
        self.top_card = top_card
        top_card.setObjectName("card")
        top_outer = QVBoxLayout(top_card)
        top_outer.setContentsMargins(10, 8, 10, 8)
        top_outer.setSpacing(6)

        file_panel = QFrame()
        file_panel.setObjectName("fileLoadPanel")
        file_layout = QHBoxLayout(file_panel)
        file_layout.setContentsMargins(0, 0, 0, 0)
        file_layout.setSpacing(12)

        self.drop_frame = DropFrame()
        self.drop_frame.file_dropped.connect(self.set_file)
        self.drop_frame.setMinimumWidth(240)
        self.drop_frame.setMaximumWidth(280)
        self.drop_frame.setMinimumHeight(70)
        self.drop_frame.setMaximumHeight(78)
        file_layout.addWidget(self.drop_frame)

        file_controls = QVBoxLayout()
        file_controls.setSpacing(6)
        file_layout.addLayout(file_controls, stretch=1)

        file_row = QHBoxLayout()
        file_row.setSpacing(10)
        self.file_label = QLineEdit("Ningun archivo seleccionado")
        self.file_label.setObjectName("filePath")
        self.file_label.setReadOnly(True)
        self.file_label.setToolTip("Ruta del archivo seleccionado")
        file_row.addWidget(self.file_label, stretch=1)

        self.select_button = QPushButton("Seleccionar archivo")
        self.select_button.setObjectName("secondaryButton")
        self.select_button.setMinimumWidth(165)
        self.select_button.clicked.connect(self.select_file)
        file_row.addWidget(self.select_button)
        file_controls.addLayout(file_row)
        # No agregar stretch vertical acá: en pantallas chicas empuja hacia abajo y achica los paneles inferiores.
        top_outer.addWidget(file_panel)

        self.performance_button = QPushButton("\u2699\ufe0f Rendimiento y transcripción")
        self.performance_button.setObjectName("performanceButton")
        self.performance_button.setMinimumHeight(42)
        self.performance_button.setToolTip("Abrir configuracion de motor, modelo, idioma, hilos, troceo y P/R")
        self.performance_button.clicked.connect(self.openPerformanceModal)
        top_outer.addWidget(self.performance_button)

        self.apply_qr_button = QPushButton("Usar P/R")
        self.apply_qr_button.setObjectName("secondaryButton")
        self.apply_qr_button.setMinimumHeight(34)
        self.apply_qr_button.setMinimumWidth(94)
        self.apply_qr_button.setToolTip("Activa formato P/R; si ya esta activo, vuelve a texto normal sin P/R")
        self.apply_qr_button.clicked.connect(self.toggle_pr_format)
        self.apply_qr_button.setEnabled(False)

        # Banner de recomendacion: concentra avisos sin afectar las opciones reales.
        self.warning_label = QLabel("\u2139\ufe0f Modo recomendado: Faster Whisper + small + Rapido + Espanol + Hilos Auto + Trocear 10 min. Para llamadas, usar P/R inteligente llamada.")
        self.warning_label.setObjectName("warning")
        self.warning_label.setWordWrap(True)
        self.warning_label.setMaximumHeight(32)
        top_outer.addWidget(self.warning_label)

        main.addWidget(top_card)

        # Barra única y compacta: antes exportación quedaba en una columna a la derecha
        # y agrandaba artificialmente toda la fila, cortando los paneles inferiores.
        actions_card = QFrame()
        self.actions_card = actions_card
        actions_card.setObjectName("actionsCard")
        actions_card.setMaximumHeight(72)
        actions_outer = QVBoxLayout(actions_card)
        actions_outer.setContentsMargins(10, 7, 10, 7)
        actions_outer.setSpacing(0)

        actions_panel = QFrame()
        actions_panel.setObjectName("actionsRow")
        actions_panel.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Minimum)
        actions = FlowLayout(actions_panel, margin=0, h_spacing=8, v_spacing=6)

        self.transcribe_button = QPushButton("\u25b6 Transcribir")
        self.transcribe_button.setObjectName("primaryButton")
        self.transcribe_button.clicked.connect(self.start_transcription)
        self.transcribe_button.setMinimumWidth(135)
        self.transcribe_button.setMinimumHeight(32)
        actions.addWidget(self.transcribe_button)

        self.cancel_button = QPushButton("\u23f9 Cancelar")
        self.cancel_button.setObjectName("cancelButton")
        self.cancel_button.clicked.connect(self.cancel_transcription)
        self.cancel_button.setEnabled(False)
        self.cancel_button.setMinimumHeight(32)
        actions.addWidget(self.cancel_button)

        self.clear_button = QPushButton("\u232b Limpiar")
        self.clear_button.setObjectName("destructiveButton")
        self.clear_button.clicked.connect(self.clear_all)
        self.clear_button.setMinimumHeight(32)
        actions.addWidget(self.clear_button)

        actions.addWidget(self.apply_qr_button)

        self.import_copy_button = QPushButton("Importar copia")
        self.import_copy_button.setObjectName("secondaryButton")
        self.import_copy_button.setMinimumHeight(32)
        self.import_copy_button.setToolTip("Abrir una copia guardada automaticamente por AR_TRANSCRIPTOR")
        self.import_copy_button.clicked.connect(self.import_saved_transcription)
        actions.addWidget(self.import_copy_button)

        export_label = QLabel("Exportar")
        export_label.setObjectName("exportLabel")
        actions.addWidget(export_label)

        self.export_txt_button = QPushButton("TXT")
        self.export_txt_button.setObjectName("chipButton")
        self.export_txt_button.clicked.connect(lambda: self.export_result("txt"))
        actions.addWidget(self.export_txt_button)

        self.export_srt_button = QPushButton("SRT")
        self.export_srt_button.setObjectName("chipButton")
        self.export_srt_button.clicked.connect(lambda: self.export_result("srt"))
        actions.addWidget(self.export_srt_button)

        self.export_vtt_button = QPushButton("VTT")
        self.export_vtt_button.setObjectName("chipButton")
        self.export_vtt_button.clicked.connect(lambda: self.export_result("vtt"))
        actions.addWidget(self.export_vtt_button)

        self.export_json_button = QPushButton("JSON")
        self.export_json_button.setObjectName("chipButton")
        self.export_json_button.clicked.connect(lambda: self.export_result("json"))
        actions.addWidget(self.export_json_button)

        self.export_docx_button = QPushButton("Word")
        self.export_docx_button.setObjectName("chipButton")
        self.export_docx_button.clicked.connect(lambda: self.export_result("docx"))
        actions.addWidget(self.export_docx_button)

        actions_outer.addWidget(actions_panel)
        main.addWidget(actions_card)
        self._set_export_enabled(False)

        workspace_card = QFrame()
        self.workspace_card = workspace_card
        workspace_card.setObjectName("workspaceCard")
        workspace_layout = QHBoxLayout(workspace_card)
        workspace_layout.setContentsMargins(10, 6, 10, 6)
        workspace_layout.setSpacing(8)
        workspace_label = QLabel("Vista")
        workspace_label.setObjectName("controlLabel")
        workspace_layout.addWidget(workspace_label)

        self.view_pro_button = QPushButton("Editor pro")
        self.view_pro_button.setObjectName("chipButton")
        self.view_pro_button.setToolTip("Vista completa: archivo, controles, progreso, texto y bloques")
        self.view_pro_button.clicked.connect(lambda: self.set_workspace_mode("pro"))
        workspace_layout.addWidget(self.view_pro_button)

        self.view_text_button = QPushButton("Solo texto")
        self.view_text_button.setObjectName("chipButton")
        self.view_text_button.setToolTip("Oculta panel derecho para editar la transcripción con más espacio")
        self.view_text_button.clicked.connect(lambda: self.set_workspace_mode("text"))
        workspace_layout.addWidget(self.view_text_button)

        self.view_focus_button = QPushButton("Enfoque")
        self.view_focus_button.setObjectName("chipButton")
        self.view_focus_button.setToolTip("Modo compacto: deja el editor y la barra de acciones, ocultando lo accesorio")
        self.view_focus_button.clicked.connect(lambda: self.set_workspace_mode("focus"))
        workspace_layout.addWidget(self.view_focus_button)

        self.view_blocks_button = QPushButton("Texto + bloques")
        self.view_blocks_button.setObjectName("chipButton")
        self.view_blocks_button.setToolTip("Muestra texto y bloques editables, con proporción equilibrada")
        self.view_blocks_button.clicked.connect(lambda: self.set_workspace_mode("blocks"))
        workspace_layout.addWidget(self.view_blocks_button)
        workspace_layout.addStretch(1)
        self.workspace_status_label = QLabel("Editor pro")
        self.workspace_status_label.setObjectName("progressMeta")
        self.workspace_status_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        workspace_layout.addWidget(self.workspace_status_label)
        main.addWidget(workspace_card)

        progress_card = QFrame()
        self.progress_card = progress_card
        progress_card.setObjectName("progressCard")
        progress_layout = QVBoxLayout(progress_card)
        progress_layout.setContentsMargins(14, 8, 14, 8)
        progress_layout.setSpacing(5)
        progress_header = QHBoxLayout()
        progress_header.setContentsMargins(0, 0, 0, 0)
        progress_header.setSpacing(8)
        self.status_label = QLabel("\u23f3 Esperando archivo...")
        self.status_label.setObjectName("status")
        self.status_label.setWordWrap(True)
        self.progress_meta_label = QLabel("0% · Duración --:--")
        self.progress_meta_label.setObjectName("progressMeta")
        self.progress_meta_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        self.progress_meta_label.setMinimumWidth(150)
        progress_header.addWidget(self.status_label, stretch=1)
        progress_header.addWidget(self.progress_meta_label)
        self.progress_bar = QProgressBar()
        self._set_progress(0)
        self.progress_bar.setTextVisible(False)
        progress_layout.addLayout(progress_header)
        progress_layout.addWidget(self.progress_bar)
        main.addWidget(progress_card)

        splitter = QSplitter(Qt.Horizontal)
        self.result_splitter = splitter
        splitter.setObjectName("resultSplitter")
        splitter.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)

        text_card = QFrame()
        self.text_card = text_card
        text_card.setObjectName("card")
        text_layout = QVBoxLayout(text_card)
        text_layout.setContentsMargins(12, 10, 12, 10)
        text_layout.setSpacing(6)
        text_header = QHBoxLayout()
        text_header.setContentsMargins(0, 0, 0, 0)
        text_header.setSpacing(8)
        text_title = QLabel("Transcripcion completa")
        text_title.setObjectName("sectionTitle")
        text_header.addWidget(text_title, stretch=1)

        self.search_input = QLineEdit()
        self.search_input.setObjectName("searchInput")
        self.search_input.setPlaceholderText("Buscar en la transcripción...")
        self.search_input.setClearButtonEnabled(True)
        self.search_input.setMaximumWidth(280)
        self.search_input.returnPressed.connect(self.search_next_match)
        self.search_input.textChanged.connect(self.search_transcription)
        text_header.addWidget(self.search_input)

        self.search_prev_button = QPushButton("‹")
        self.search_prev_button.setObjectName("chipButton")
        self.search_prev_button.setToolTip("Resultado anterior")
        self.search_prev_button.setFixedWidth(34)
        self.search_prev_button.clicked.connect(self.search_previous_match)
        text_header.addWidget(self.search_prev_button)

        self.search_next_button = QPushButton("›")
        self.search_next_button.setObjectName("chipButton")
        self.search_next_button.setToolTip("Resultado siguiente")
        self.search_next_button.setFixedWidth(34)
        self.search_next_button.clicked.connect(self.search_next_match)
        text_header.addWidget(self.search_next_button)

        self.search_status_label = QLabel("0")
        self.search_status_label.setObjectName("progressMeta")
        self.search_status_label.setMinimumWidth(54)
        self.search_status_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        text_header.addWidget(self.search_status_label)

        self.text_edit = QTextEdit()
        self.text_edit.setPlaceholderText("Aca va a aparecer el texto transcripto...")
        self.text_edit.textChanged.connect(self._on_transcription_text_changed)
        text_layout.addLayout(text_header)
        text_layout.addWidget(self.text_edit, stretch=1)

        seg_card = QFrame()
        self.seg_card = seg_card
        seg_card.setObjectName("prPanel")
        seg_card.setMinimumWidth(360)
        seg_card.setMinimumHeight(220)
        seg_layout = QVBoxLayout(seg_card)
        seg_layout.setContentsMargins(12, 10, 12, 10)
        seg_layout.setSpacing(6)

        # Encabezado siempre visible: antes el botón quedaba abajo y se cortaba en pantallas bajas.
        pr_header = QHBoxLayout()
        pr_header.setContentsMargins(0, 0, 0, 0)
        pr_header.setSpacing(8)
        seg_title = QLabel("Bloques P/R editables")
        seg_title.setObjectName("sectionTitle")
        pr_header.addWidget(seg_title, stretch=1)

        self.edit_pr_block_top_button = QPushButton("✎ Editar bloque")
        self.edit_pr_block_top_button.setObjectName("primaryButton")
        self.edit_pr_block_top_button.setToolTip("Abre el modal para modificar etiqueta/interlocutor y texto del bloque seleccionado")
        self.edit_pr_block_top_button.setMinimumHeight(32)
        self.edit_pr_block_top_button.clicked.connect(lambda: self.openEditBlockModal())
        pr_header.addWidget(self.edit_pr_block_top_button)

        self.speakers_button = QPushButton("👥 Interlocutores")
        self.speakers_button.setObjectName("secondaryButton")
        self.speakers_button.setToolTip("Detectar/renombrar Speaker 1 y Speaker 2 y aplicar esos nombres a toda la transcripción")
        self.speakers_button.setMinimumHeight(32)
        self.speakers_button.clicked.connect(self.open_speakers_modal)
        pr_header.addWidget(self.speakers_button)

        self.review_button = QPushButton("🔎 Revisar")
        self.review_button.setObjectName("secondaryButton")
        self.review_button.setToolTip("Abrir revisión inteligente: palabras clave, silencios, segmentos largos y baja confianza")
        self.review_button.setMinimumHeight(32)
        self.review_button.clicked.connect(self.open_review_modal)
        pr_header.addWidget(self.review_button)

        pr_help = QLabel("Etiquetas editables: Speaker 1/Speaker 2, PREGUNTA/RESPUESTA u otros nombres. Elegí un bloque y tocá Editar.")
        pr_help.setObjectName("miniHelp")
        pr_help.setWordWrap(True)

        pr_names_frame = QFrame()
        pr_names_frame.setObjectName("prLabelPanel")
        pr_names = QGridLayout(pr_names_frame)
        pr_names.setContentsMargins(0, 0, 0, 0)
        pr_names.setHorizontalSpacing(8)
        pr_names.setVerticalSpacing(8)
        pr_names.setColumnStretch(0, 0)
        pr_names.setColumnStretch(1, 1)
        pr_names.setColumnStretch(2, 0)
        pr_names.setColumnStretch(3, 1)
        qn = QLabel("Etiqueta 1")
        qn.setObjectName("controlLabel")
        self.question_label_edit = QLineEdit("PREGUNTA")
        self.question_label_edit.setObjectName("smallInput")
        self.question_label_edit.setPlaceholderText("Ej.: PREGUNTA / OPERADOR")
        self.question_label_edit.setToolTip("Nombre que queres usar para el primer rol. Ej.: OPERADOR o LLAMANTE")
        self.question_label_edit.textChanged.connect(self._pr_label_names_changed)
        an = QLabel("Etiqueta 2")
        an.setObjectName("controlLabel")
        self.answer_label_edit = QLineEdit("RESPUESTA")
        self.answer_label_edit.setObjectName("smallInput")
        self.answer_label_edit.setPlaceholderText("Ej.: RESPUESTA / LLAMANTE")
        self.answer_label_edit.setToolTip("Nombre que queres usar para el segundo rol. Ej.: LLAMANTE, REQUIRIENTE o ENTREVISTADO")
        self.answer_label_edit.textChanged.connect(self._pr_label_names_changed)
        pr_names.addWidget(qn, 0, 0)
        pr_names.addWidget(self.question_label_edit, 0, 1)
        pr_names.addWidget(an, 0, 2)
        pr_names.addWidget(self.answer_label_edit, 0, 3)

        self.segments_list = QListWidget()
        self.segments_list.setObjectName("prBlockList")
        self.segments_list.setMinimumHeight(55)
        self.segments_list.setMaximumHeight(9999)
        self.segments_list.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.segments_list.itemClicked.connect(lambda _item: self.openEditBlockModal())
        self.segments_list.currentRowChanged.connect(self._pr_selection_changed)

        pr_buttons_frame = QFrame()
        pr_buttons_frame.setObjectName("prActions")
        pr_buttons = FlowLayout(pr_buttons_frame, margin=0, h_spacing=8, v_spacing=8)
        self.edit_pr_block_button = QPushButton("✎ Editar")
        self.edit_pr_block_button.setObjectName("secondaryButton")
        self.edit_pr_block_button.setToolTip("Abre un editor para modificar la etiqueta y el texto del bloque seleccionado")
        self.edit_pr_block_button.clicked.connect(lambda: self.openEditBlockModal())
        self.edit_pr_block_button.setMinimumHeight(32)
        pr_buttons.addWidget(self.edit_pr_block_button)

        self.mark_question_button = QPushButton("\U0001F3F7\ufe0f Marcar etiqueta 1")
        self.mark_question_button.setObjectName("secondaryButton")
        self.mark_question_button.setToolTip("Marca el bloque seleccionado con la etiqueta 1")
        self.mark_question_button.clicked.connect(self.marcarEtiqueta1)
        self.mark_question_button.setMinimumHeight(32)
        pr_buttons.addWidget(self.mark_question_button)

        self.mark_answer_button = QPushButton("\U0001F3F7\ufe0f Marcar etiqueta 2")
        self.mark_answer_button.setObjectName("secondaryButton")
        self.mark_answer_button.setToolTip("Marca el bloque seleccionado con la etiqueta 2")
        self.mark_answer_button.clicked.connect(self.marcarEtiqueta2)
        self.mark_answer_button.setMinimumHeight(32)
        pr_buttons.addWidget(self.mark_answer_button)

        self.toggle_pr_button = QPushButton("🔄 Alternar")
        self.toggle_pr_button.setObjectName("secondaryButton")
        self.toggle_pr_button.setToolTip("Cambia PREGUNTA por RESPUESTA o al reves")
        self.toggle_pr_button.clicked.connect(self.alternarEtiquetas)
        self.toggle_pr_button.setMinimumHeight(32)
        pr_buttons.addWidget(self.toggle_pr_button)

        self.regen_pr_button = QPushButton("\U0000267B\ufe0f Regenerar P/R")
        self.regen_pr_button.setObjectName("secondaryButton")
        self.regen_pr_button.setToolTip("Vuelve a generar P/R desde los segmentos originales")
        self.regen_pr_button.clicked.connect(self.regenerarPR)
        self.regen_pr_button.setMinimumHeight(32)
        pr_buttons.addWidget(self.regen_pr_button)

        seg_layout.addLayout(pr_header)
        seg_layout.addWidget(pr_help)
        seg_layout.addWidget(pr_names_frame)
        # Botonera arriba de la lista: así no queda cortada en pantallas bajas.
        seg_layout.addWidget(pr_buttons_frame)
        seg_layout.addWidget(self.segments_list, stretch=1)

        splitter.addWidget(text_card)
        splitter.addWidget(seg_card)
        splitter.setSizes([900, 390])
        main.addWidget(splitter, stretch=1)

        footer_row = QHBoxLayout()
        footer = QLabel("Tip: Ctrl+F buscar · Ctrl+R revisar · Ctrl+S guardar · F2 editar bloque · Ctrl+1/2/3 cambiar vista.")
        footer.setObjectName("footer")
        footer.setWordWrap(True)
        # Footer: ayuda breve y acceso directo de donacion por MercadoPago.
        self.donate_button = QPushButton("\u2615 Invitame un caf\u00e9")
        self.donate_button.setObjectName("donateButton")
        self.donate_button.setToolTip("Apoy\u00e1 el desarrollo de esta herramienta")
        self.donate_button.clicked.connect(self.open_donation)
        footer_row.addWidget(footer, stretch=1)
        footer_row.addWidget(self.donate_button)
        footer.setMaximumHeight(28)
        self.donate_button.setMaximumHeight(34)
        main.addLayout(footer_row)
        self._set_pr_edit_enabled(False)
        self._setup_shortcuts()
        self.set_workspace_mode("pro")

    def set_workspace_mode(self, mode: str) -> None:
        """Cambia la disposición visual sin tocar la transcripción ni los datos."""
        mode = mode if mode in {"pro", "text", "focus", "blocks"} else "pro"
        self.workspace_mode = mode

        # Valores base: vista completa.
        if hasattr(self, "top_card"):
            self.top_card.setVisible(mode != "focus")
        if hasattr(self, "progress_card"):
            self.progress_card.setVisible(mode != "focus")
        if hasattr(self, "workspace_card"):
            self.workspace_card.setVisible(True)
        if hasattr(self, "seg_card"):
            self.seg_card.setVisible(mode in {"pro", "blocks"})

        if hasattr(self, "result_splitter"):
            if mode == "text":
                self.result_splitter.setSizes([1200, 0])
            elif mode == "focus":
                self.result_splitter.setSizes([1200, 0])
            elif mode == "blocks":
                self.result_splitter.setSizes([760, 520])
            else:
                self.result_splitter.setSizes([900, 390])

        labels = {
            "pro": "Editor pro",
            "text": "Solo texto",
            "focus": "Enfoque",
            "blocks": "Texto + bloques",
        }
        if hasattr(self, "workspace_status_label"):
            self.workspace_status_label.setText(labels.get(mode, "Editor pro"))
        if hasattr(self, "status_label"):
            self.status_label.setText(f"Vista activa: {labels.get(mode, 'Editor pro')}.")

    def _apply_style(self) -> None:
        self.setStyleSheet(build_stylesheet(self.theme_name))
        self.theme_toggle_button.setText(theme_toggle_text(self.theme_name))

    def toggle_theme(self) -> None:
        self.theme_name = "light" if self.theme_name == "dark" else "dark"
        save_theme_name(self.theme_name)
        self._apply_style()

    def open_donation(self) -> None:
        webbrowser.open("https://mpago.la/2d7UFxJ")

    def _setup_shortcuts(self) -> None:
        shortcuts = [
            ("Ctrl+F", lambda: self._focus_search()),
            ("Ctrl+S", lambda: self._autosave_transcription("atajo_ctrl_s")),
            ("Ctrl+O", self.select_file),
            ("Ctrl+Return", self.start_transcription),
            ("F2", lambda: self.openEditBlockModal()),
            ("F3", self.search_next_match),
            ("Ctrl+R", self.open_review_modal),
            ("Shift+F3", self.search_previous_match),
            ("Ctrl+1", lambda: self.set_workspace_mode("pro")),
            ("Ctrl+2", lambda: self.set_workspace_mode("text")),
            ("Ctrl+3", lambda: self.set_workspace_mode("focus")),
            ("Ctrl+4", lambda: self.set_workspace_mode("blocks")),
        ]
        for key, callback in shortcuts:
            shortcut = QShortcut(QKeySequence(key), self)
            shortcut.activated.connect(callback)

    def _focus_search(self) -> None:
        if hasattr(self, "search_input"):
            self.search_input.setFocus()
            self.search_input.selectAll()


    def open_review_modal(self) -> None:
        """Panel general de revisión: no modifica el texto; ayuda a encontrar partes para controlar."""
        dialog = QDialog(self)
        dialog.setWindowTitle("Revisión inteligente")
        dialog.setModal(True)
        dialog.resize(820, 560)
        layout = QVBoxLayout(dialog)
        layout.setContentsMargins(16, 14, 16, 14)
        layout.setSpacing(10)

        title = QLabel("🔎 Revisión inteligente")
        title.setObjectName("sectionTitle")
        layout.addWidget(title)

        help_label = QLabel("Detecta palabras clave, silencios, segmentos largos y posibles palabras de baja confianza. Doble click en un resultado para ubicarlo.")
        help_label.setObjectName("miniHelp")
        help_label.setWordWrap(True)
        layout.addWidget(help_label)

        keyword_row = QHBoxLayout()
        keyword_label = QLabel("Palabras clave")
        keyword_label.setObjectName("controlLabel")
        keyword_row.addWidget(keyword_label)
        keywords_edit = QLineEdit(", ".join(self.review_keywords))
        keywords_edit.setPlaceholderText("Ej.: nombre, dirección, patente, dinero, reunión")
        keywords_edit.setToolTip("Separá palabras o frases con coma")
        keyword_row.addWidget(keywords_edit, stretch=1)
        analyze_button = QPushButton("Analizar")
        analyze_button.setObjectName("primaryButton")
        keyword_row.addWidget(analyze_button)
        layout.addLayout(keyword_row)

        findings_list = QListWidget()
        findings_list.setObjectName("prBlockList")
        layout.addWidget(findings_list, stretch=1)

        summary_label = QLabel("Sin análisis todavía.")
        summary_label.setObjectName("progressMeta")
        layout.addWidget(summary_label)

        actions = QHBoxLayout()
        actions.addStretch(1)
        close_button = QPushButton("Cerrar")
        close_button.setObjectName("secondaryButton")
        close_button.clicked.connect(dialog.accept)
        actions.addWidget(close_button)
        layout.addLayout(actions)

        def refresh() -> None:
            keywords = [x.strip() for x in keywords_edit.text().split(",") if x.strip()]
            self.review_keywords = keywords
            findings = self._collect_review_findings(keywords)
            self.review_findings = findings
            findings_list.clear()
            for finding in findings:
                item = QListWidgetItem(self._format_review_finding(finding))
                item.setData(Qt.UserRole, finding)
                findings_list.addItem(item)
            summary_label.setText(self._review_summary(findings))

        def open_selected(item: QListWidgetItem) -> None:
            finding = item.data(Qt.UserRole)
            if isinstance(finding, dict):
                self._jump_to_review_finding(finding)
                dialog.accept()

        analyze_button.clicked.connect(refresh)
        findings_list.itemDoubleClicked.connect(open_selected)
        refresh()
        self._center_dialog(dialog)
        dialog.exec()

    def _collect_review_findings(self, keywords: List[str]) -> List[Dict[str, Any]]:
        findings: List[Dict[str, Any]] = []
        segments = (self.last_result or {}).get("segments", []) if self.last_result else []
        if not isinstance(segments, list):
            segments = []
        blocks = self.pr_blocks if self.pr_blocks else [seg for seg in segments if isinstance(seg, dict)]

        def clean(value: str) -> str:
            return _remove_accents_for_match(str(value or "")).lower()

        # Palabras clave configuradas por el usuario.
        for term in keywords:
            term_norm = clean(term)
            if not term_norm:
                continue
            for idx, block in enumerate(blocks):
                text = str(block.get("text", ""))
                if term_norm in clean(text):
                    findings.append({
                        "kind": "keyword",
                        "title": f"Palabra clave: {term}",
                        "detail": text,
                        "index": idx,
                        "start": float(block.get("start", 0) or 0),
                        "end": float(block.get("end", 0) or 0),
                        "term": term,
                    })

        # Segmentos largos: suelen convenir para dividir o revisar puntuación.
        for idx, block in enumerate(blocks):
            start = float(block.get("start", 0) or 0)
            end = float(block.get("end", start) or start)
            duration = max(0.0, end - start)
            if duration >= 24:
                findings.append({
                    "kind": "long_segment",
                    "title": f"Segmento largo: {seconds_to_duration_label(duration)}",
                    "detail": str(block.get("text", "")),
                    "index": idx,
                    "start": start,
                    "end": end,
                })

        # Silencios/gaps entre segmentos originales.
        prev_end: Optional[float] = None
        for idx, seg in enumerate(segments):
            if not isinstance(seg, dict):
                continue
            start = float(seg.get("start", 0) or 0)
            end = float(seg.get("end", start) or start)
            if prev_end is not None:
                gap = start - prev_end
                if gap >= 3.0:
                    findings.append({
                        "kind": "silence",
                        "title": f"Silencio detectado: {seconds_to_duration_label(gap)}",
                        "detail": "Pausa prolongada entre fragmentos.",
                        "index": idx,
                        "start": prev_end,
                        "end": start,
                    })
            prev_end = max(prev_end or 0.0, end)

        # Baja confianza por palabra cuando Faster-Whisper entrega probability.
        for idx, seg in enumerate(segments):
            if not isinstance(seg, dict):
                continue
            low_words = []
            for w in seg.get("words", []) or []:
                if not isinstance(w, dict) or "probability" not in w:
                    continue
                try:
                    prob = float(w.get("probability"))
                except (TypeError, ValueError):
                    continue
                if prob < 0.55:
                    low_words.append(str(w.get("word", "")).strip())
            if low_words:
                sample = ", ".join([x for x in low_words[:8] if x])
                findings.append({
                    "kind": "low_confidence",
                    "title": "Posible baja confianza",
                    "detail": f"Palabras a revisar: {sample}" if sample else str(seg.get("text", "")),
                    "index": idx,
                    "start": float(seg.get("start", 0) or 0),
                    "end": float(seg.get("end", 0) or 0),
                    "term": low_words[0] if low_words else "",
                })

        findings.sort(key=lambda x: (float(x.get("start", 0) or 0), str(x.get("kind", ""))))
        return findings

    def _format_review_finding(self, finding: Dict[str, Any]) -> str:
        start = seconds_to_hhmmss(float(finding.get("start", 0) or 0))
        end = seconds_to_hhmmss(float(finding.get("end", 0) or 0))
        detail = re.sub(r"\s+", " ", str(finding.get("detail", "")).strip())
        if len(detail) > 135:
            detail = detail[:132] + "..."
        icon = {
            "keyword": "🔑",
            "silence": "▫",
            "long_segment": "↔",
            "low_confidence": "⚠",
        }.get(str(finding.get("kind", "")), "•")
        return f"{icon} [{start} - {end}] {finding.get('title', 'Revisar')} · {detail}"

    def _review_summary(self, findings: List[Dict[str, Any]]) -> str:
        counts: Dict[str, int] = {}
        for f in findings:
            key = str(f.get("kind", "otro"))
            counts[key] = counts.get(key, 0) + 1
        if not counts:
            return "Sin hallazgos para revisar."
        labels = {
            "keyword": "palabras clave",
            "silence": "silencios",
            "long_segment": "segmentos largos",
            "low_confidence": "baja confianza",
        }
        parts = [f"{count} {labels.get(kind, kind)}" for kind, count in counts.items()]
        return "Hallazgos: " + " · ".join(parts)

    def _jump_to_review_finding(self, finding: Dict[str, Any]) -> None:
        index = int(finding.get("index", 0) or 0)
        if hasattr(self, "segments_list") and self.segments_list.count():
            self.segments_list.setCurrentRow(max(0, min(index, self.segments_list.count() - 1)))
        term = str(finding.get("term", "")).strip()
        if term and hasattr(self, "search_input"):
            self.search_input.setText(term)
            self.search_transcription(term)
        if hasattr(self, "status_label"):
            self.status_label.setText(f"Revisión ubicada en {seconds_to_hhmmss(float(finding.get('start', 0) or 0))}: {finding.get('title', 'Revisar')}.")

    def _on_transcription_text_changed(self) -> None:
        # En importaciones grandes no recalculamos búsqueda/highlights en cada setPlainText,
        # porque puede congelar la ventana con textos largos.
        if getattr(self, "_bulk_loading_ui", False):
            return
        if hasattr(self, "search_input") and self.search_input.text().strip():
            self.search_transcription(self.search_input.text())

    def _clear_search_highlights(self) -> None:
        if hasattr(self, "text_edit"):
            self.text_edit.setExtraSelections([])

    def search_transcription(self, query: Optional[str] = None) -> None:
        if not hasattr(self, "text_edit"):
            return
        query = self.search_input.text() if query is None and hasattr(self, "search_input") else str(query or "")
        query = query.strip()
        self.search_matches = []
        self.search_match_index = -1
        if not query:
            self._clear_search_highlights()
            if hasattr(self, "search_status_label"):
                self.search_status_label.setText("0")
            return
        text = self.text_edit.toPlainText()
        pattern = re.escape(query)
        flags = re.IGNORECASE | re.UNICODE
        self.search_matches = []
        for m in re.finditer(pattern, text, flags):
            self.search_matches.append((m.start(), m.end()))
            # Evita miles de ExtraSelections que hacen pesada la app.
            if len(self.search_matches) >= MAX_SEARCH_HIGHLIGHTS:
                break
        if self.search_matches:
            self.search_match_index = 0
        self._paint_search_matches()
        if self.search_matches:
            self._go_to_search_match(0)
            if len(self.search_matches) >= MAX_SEARCH_HIGHLIGHTS:
                self.status_label.setText(f"Búsqueda limitada a los primeros {MAX_SEARCH_HIGHLIGHTS} resultados para mantener fluida la app.")
        else:
            self.status_label.setText(f"Sin resultados para: {query}")

    def _paint_search_matches(self) -> None:
        selections = []
        text_doc = self.text_edit.document()
        for idx, (start, end) in enumerate(self.search_matches):
            cursor = QTextCursor(text_doc)
            cursor.setPosition(start)
            cursor.setPosition(end, QTextCursor.KeepAnchor)
            fmt = QTextCharFormat()
            if idx == self.search_match_index:
                fmt.setBackground(QColor(251, 146, 60, 180))
                fmt.setForeground(QColor(255, 255, 255))
            else:
                fmt.setBackground(QColor(251, 191, 36, 90))
            sel = QTextEdit.ExtraSelection()
            sel.cursor = cursor
            sel.format = fmt
            selections.append(sel)
        self.text_edit.setExtraSelections(selections)
        if hasattr(self, "search_status_label"):
            total = len(self.search_matches)
            current = self.search_match_index + 1 if total and self.search_match_index >= 0 else 0
            self.search_status_label.setText(f"{current}/{total}" if total else "0")

    def _go_to_search_match(self, index: int) -> None:
        if not self.search_matches:
            self._paint_search_matches()
            return
        self.search_match_index = index % len(self.search_matches)
        start, end = self.search_matches[self.search_match_index]
        cursor = self.text_edit.textCursor()
        cursor.setPosition(start)
        cursor.setPosition(end, QTextCursor.KeepAnchor)
        self.text_edit.setTextCursor(cursor)
        self.text_edit.ensureCursorVisible()
        self._paint_search_matches()
        query = self.search_input.text().strip() if hasattr(self, "search_input") else ""
        self.status_label.setText(f"Búsqueda: {self.search_match_index + 1}/{len(self.search_matches)} para '{query}'.")

    def search_next_match(self) -> None:
        if not hasattr(self, "search_input"):
            return
        if not self.search_matches and self.search_input.text().strip():
            self.search_transcription(self.search_input.text())
        if self.search_matches:
            self._go_to_search_match(self.search_match_index + 1)

    def search_previous_match(self) -> None:
        if not hasattr(self, "search_input"):
            return
        if not self.search_matches and self.search_input.text().strip():
            self.search_transcription(self.search_input.text())
        if self.search_matches:
            self._go_to_search_match(self.search_match_index - 1)

    def _check_environment(self) -> None:
        if shutil.which("ffmpeg") is None:
            self.status_label.setText("Atencion: FFmpeg no esta instalado o no esta en PATH.")
            QMessageBox.warning(
                self,
                "FFmpeg requerido",
                "FFmpeg no esta instalado o no esta en PATH.\n\n"
                "Whisper lo necesita para procesar audios y videos.\n"
                "Instalalo antes de transcribir archivos."
            )

    def _set_progress(self, value: int) -> None:
        value = max(0, min(100, int(value or 0)))
        self.progress_bar.setValue(value)
        self._refresh_progress_meta(value)

    def _refresh_progress_meta(self, value: Optional[int] = None) -> None:
        if not hasattr(self, "progress_meta_label"):
            return
        current = self.progress_bar.value() if value is None and hasattr(self, "progress_bar") else int(value or 0)
        duration = seconds_to_duration_label(self.media_duration_seconds)
        self.progress_meta_label.setText(f"{current}% · Duración {duration}")
        tooltip = "Porcentaje de avance de la transcripción y duración total del audio/video cargado."
        if self.media_duration_seconds:
            tooltip += f" Duración detectada: {seconds_to_hhmmss(self.media_duration_seconds)}."
        self.progress_meta_label.setToolTip(tooltip)

    def _detect_selected_file_duration(self) -> None:
        self.media_duration_seconds = probe_media_duration_seconds(self.selected_file or "")
        self._refresh_progress_meta()

    def openPerformanceModal(self, *_args: Any) -> None:
        print("CLICK EN RENDIMIENTO Y TRANSCRIPCION")
        self.open_performance_settings()

    def _center_dialog(self, dialog: QDialog) -> None:
        parent_center = self.frameGeometry().center()
        dialog.move(parent_center - dialog.rect().center())

    def open_performance_settings(self) -> None:
        dialog = QDialog(self)
        dialog.setObjectName("settingsDialog")
        dialog.setWindowTitle("\u2699\ufe0f Rendimiento y transcripción")
        dialog.setModal(True)
        dialog.resize(760, 560)

        layout = QVBoxLayout(dialog)
        layout.setContentsMargins(18, 18, 18, 18)
        layout.setSpacing(12)

        title = QLabel("\u2699\ufe0f Rendimiento y transcripción")
        title.setObjectName("sectionTitle")
        layout.addWidget(title)

        grid_frame = QFrame()
        grid_frame.setObjectName("modalGrid")
        grid = QGridLayout(grid_frame)
        grid.setContentsMargins(0, 0, 0, 0)
        grid.setHorizontalSpacing(12)
        grid.setVerticalSpacing(12)
        for col in range(3):
            grid.setColumnStretch(col, 1)

        def add_modal_field(row: int, col: int, label_text: str, widget: QWidget) -> None:
            label = QLabel(label_text.upper())
            label.setObjectName("controlLabel")
            box = QFrame()
            box.setObjectName("fieldGroup")
            box_layout = QVBoxLayout(box)
            box_layout.setContentsMargins(0, 0, 0, 0)
            box_layout.setSpacing(5)
            box_layout.addWidget(label)
            box_layout.addWidget(widget)
            widget.setMinimumWidth(190)
            widget.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
            grid.addWidget(box, row, col)

        engine_combo = QComboBox()
        engine_combo.addItems(list(ENGINE_OPTIONS.keys()))
        engine_combo.setCurrentText(str(self.settings.get("engine", "Faster Whisper (mas rapido)")))
        add_modal_field(0, 0, "Motor", engine_combo)

        model_combo = QComboBox()
        model_combo.addItems(MODELS)
        model_combo.setCurrentText(str(self.settings.get("model", "small")))
        add_modal_field(0, 1, "Modelo", model_combo)

        language_combo = QComboBox()
        language_combo.addItems(LANGUAGES.keys())
        language_combo.setCurrentText(str(self.settings.get("language", "Espanol")))
        add_modal_field(0, 2, "Idioma", language_combo)

        performance_combo = QComboBox()
        performance_combo.addItems(PERFORMANCE_PROFILES.keys())
        performance_combo.setCurrentText(str(self.settings.get("performance", "Rapido")))
        add_modal_field(1, 0, "Modo de rendimiento", performance_combo)

        thread_combo = QComboBox()
        thread_combo.addItems(THREAD_OPTIONS)
        thread_combo.setCurrentText(str(self.settings.get("threads", "Auto")))
        add_modal_field(1, 1, "Hilos", thread_combo)

        chunk_combo = QComboBox()
        chunk_combo.addItems(list(CHUNK_OPTIONS.keys()))
        chunk_combo.setCurrentText(str(self.settings.get("chunk", "10 min")))
        add_modal_field(1, 2, "Troceo", chunk_combo)

        compute_combo = QComboBox()
        compute_combo.addItems(list(COMPUTE_OPTIONS.keys()))
        compute_combo.setCurrentText(str(self.settings.get("compute", "Auto")))
        add_modal_field(2, 0, "INT8 / FP16", compute_combo)

        device_combo = QComboBox()
        device_combo.addItems(list(DEVICE_OPTIONS.keys()))
        device_combo.setCurrentText(str(self.settings.get("device", "Automatico")))
        add_modal_field(2, 1, "CPU / GPU", device_combo)

        task_combo = QComboBox()
        task_combo.addItems(TASKS.keys())
        task_combo.setCurrentText(str(self.settings.get("task", "Transcribir")))
        add_modal_field(2, 2, "Tarea", task_combo)

        format_combo = QComboBox()
        format_combo.addItems(FORMAT_MODES.keys())
        format_combo.setCurrentText(str(self.settings.get("format", "Texto normal")))
        add_modal_field(3, 0, "Modo llamada / P-R", format_combo)

        manual_language_combo = QComboBox()
        manual_language_combo.setEditable(True)
        manual_language_combo.addItems(["", "es", "en", "pt", "it", "fr", "de", "ja", "zh", "ko", "ru", "ar"])
        manual_language_combo.setCurrentText(str(self.settings.get("manual_language", "")))
        manual_language_combo.setEnabled(LANGUAGES.get(language_combo.currentText()) == "manual")
        language_combo.currentTextChanged.connect(lambda text: manual_language_combo.setEnabled(LANGUAGES.get(text) == "manual"))
        add_modal_field(3, 1, "Idioma manual", manual_language_combo)

        optimize_audio_check = QCheckBox("Normalizar audio")
        optimize_audio_check.setChecked(bool(self.settings.get("optimize_audio", True)))
        grid.addWidget(optimize_audio_check, 4, 0)

        clean_text_check = QCheckBox("Limpieza básica")
        clean_text_check.setChecked(bool(self.settings.get("clean_text", True)))
        grid.addWidget(clean_text_check, 4, 1)

        pr_editable_check = QCheckBox("P/R editable activado")
        pr_editable_check.setChecked(FORMAT_MODES.get(format_combo.currentText(), "normal") in {"qr_auto", "qr_call"})
        grid.addWidget(pr_editable_check, 4, 2)

        smart_call_check = QCheckBox("P/R inteligente llamada")
        smart_call_check.setChecked(FORMAT_MODES.get(format_combo.currentText(), "normal") == "qr_call")
        grid.addWidget(smart_call_check, 5, 0)

        split_pr_check = QCheckBox("P/R alternado / separar pregunta-respuesta")
        split_pr_check.setChecked(FORMAT_MODES.get(format_combo.currentText(), "normal") == "qr_auto")
        grid.addWidget(split_pr_check, 5, 1)

        syncing = {"active": False}

        def sync_checks_from_format(*_args: Any) -> None:
            syncing["active"] = True
            mode = FORMAT_MODES.get(format_combo.currentText(), "normal")
            pr_editable_check.setChecked(mode in {"qr_auto", "qr_call"})
            smart_call_check.setChecked(mode == "qr_call")
            split_pr_check.setChecked(mode == "qr_auto")
            syncing["active"] = False

        def pr_editable_changed(*_args: Any) -> None:
            if syncing["active"]:
                return
            if not pr_editable_check.isChecked():
                format_combo.setCurrentText("Texto normal")
            elif FORMAT_MODES.get(format_combo.currentText(), "normal") == "normal":
                format_combo.setCurrentText("P/R inteligente llamada")

        def smart_changed(*_args: Any) -> None:
            if syncing["active"]:
                return
            format_combo.setCurrentText("P/R inteligente llamada" if smart_call_check.isChecked() else "Texto normal")

        def split_changed(*_args: Any) -> None:
            if syncing["active"]:
                return
            format_combo.setCurrentText("P/R alternado simple" if split_pr_check.isChecked() else "Texto normal")

        format_combo.currentTextChanged.connect(sync_checks_from_format)
        pr_editable_check.stateChanged.connect(pr_editable_changed)
        smart_call_check.stateChanged.connect(smart_changed)
        split_pr_check.stateChanged.connect(split_changed)

        layout.addWidget(grid_frame)

        actions_frame = QFrame()
        actions_frame.setObjectName("modalActions")
        actions = FlowLayout(actions_frame, margin=0, h_spacing=8, v_spacing=8)
        cancel_button = QPushButton("Cancelar")
        cancel_button.setObjectName("cancelButton")
        cancel_button.setMinimumHeight(38)
        cancel_button.clicked.connect(lambda: self.closePerformanceModal(dialog))
        save_button = QPushButton("Guardar configuracion")
        save_button.setObjectName("primaryButton")
        save_button.setMinimumHeight(38)

        controls = {
            "engine": engine_combo,
            "model": model_combo,
            "language": language_combo,
            "manual_language": manual_language_combo,
            "performance": performance_combo,
            "threads": thread_combo,
            "chunk": chunk_combo,
            "compute": compute_combo,
            "device": device_combo,
            "task": task_combo,
            "format": format_combo,
            "optimize_audio": optimize_audio_check,
            "clean_text": clean_text_check,
        }
        save_button.clicked.connect(lambda: self.savePerformanceSettings(dialog, controls))
        actions.addWidget(cancel_button)
        actions.addWidget(save_button)
        layout.addWidget(actions_frame)
        self._center_dialog(dialog)
        dialog.exec()

    def closePerformanceModal(self, dialog: QDialog) -> None:
        dialog.reject()

    def savePerformanceSettings(self, dialog: QDialog, controls: Dict[str, QWidget]) -> None:
        self.settings.update({
            "engine": controls["engine"].currentText(),
            "model": controls["model"].currentText(),
            "language": controls["language"].currentText(),
            "manual_language": controls["manual_language"].currentText().strip(),
            "performance": controls["performance"].currentText(),
            "threads": controls["threads"].currentText(),
            "chunk": controls["chunk"].currentText(),
            "compute": controls["compute"].currentText(),
            "device": controls["device"].currentText(),
            "task": controls["task"].currentText(),
            "format": controls["format"].currentText(),
            "optimize_audio": controls["optimize_audio"].isChecked(),
            "clean_text": controls["clean_text"].isChecked(),
        })
        self._validate_options()
        if FORMAT_MODES.get(str(self.settings.get("format", "Texto normal")), "normal") == "normal":
            self.pr_blocks = []
            self.pr_notes = []
            if self.last_result:
                text = str(self.last_result.get("text", "")).strip()
                segments = self.last_result.get("segments", []) or []
                self.text_edit.setPlainText(build_formatted_text("normal", segments, text, cleanup=bool(self.settings.get("clean_text", True))))
                self._show_raw_segments(segments)
        self._update_pr_format_button()
        self.status_label.setText("Configuracion de rendimiento actualizada.")
        dialog.accept()

    def _validate_options(self) -> None:
        model = str(self.settings.get("model", "small"))
        task = TASKS.get(str(self.settings.get("task", "Transcribir")))
        profile = PERFORMANCE_PROFILES.get(str(self.settings.get("performance", "Rapido")), "fast")
        warnings = []
        if model == "turbo" and task == "translate":
            warnings.append("Aviso: turbo esta pensado para transcribir; para traducir conviene medium o large.")
        mode = FORMAT_MODES.get(str(self.settings.get("format", "Texto normal")), "normal")
        if mode == "qr_auto":
            warnings.append("P/R alternado es solo borrador: puede mezclar voces si Whisper corta frases cortas.")
        elif mode == "qr_call":
            warnings.append("P/R inteligente llamada agrupa y estima operador/requirente, pero no reemplaza la revision del audio.")
        if bool(self.settings.get("clean_text", True)):
            warnings.append("Limpieza basica activa: corrige frases comunes y reduce repeticiones sospechosas.")
        engine = ENGINE_OPTIONS.get(str(self.settings.get("engine", "Faster Whisper (mas rapido)")), "faster")
        if engine == "faster":
            warnings.append("Motor rapido activo: Faster Whisper suele mejorar mucho la velocidad en CPU.")
        else:
            warnings.append("Motor oficial OpenAI: mas compatible, pero normalmente mas lento.")
        if COMPUTE_OPTIONS.get(str(self.settings.get("compute", "Auto")), "auto") == "int8":
            warnings.append("int8 CPU rapido: recomendado para computadoras sin placa de video NVIDIA.")
        if profile == "fast":
            warnings.append("Rapido: menor demora, puede bajar algo la precision.")
        elif profile == "accurate":
            warnings.append("Preciso: mejor resultado, mas lento.")
        else:
            warnings.append("Equilibrado: buen punto medio.")
        chunk_seconds = CHUNK_OPTIONS.get(str(self.settings.get("chunk", "10 min")), 0)
        if chunk_seconds:
            warnings.append(f"Troceo activo cada {chunk_seconds // 60} min: mejor para archivos largos y cancelacion.")
        self.warning_label.setText("\u2139\ufe0f " + " ".join(warnings))

    def _format_changed(self) -> None:
        self._validate_options()

    def _sync_format_checks(self) -> None:
        return

    def _smart_call_check_changed(self) -> None:
        return

    def _split_pr_check_changed(self) -> None:
        return


    def _is_pr_mode(self) -> bool:
        mode = FORMAT_MODES.get(str(self.settings.get("format", "Texto normal")), "normal")
        return mode in {"qr_auto", "qr_call"}

    def _current_pr_note_flag(self) -> bool:
        mode = FORMAT_MODES.get(str(self.settings.get("format", "Texto normal")), "normal")
        return mode == "qr_call"

    def _question_label(self) -> str:
        value = self.question_label_edit.text().strip() if hasattr(self, "question_label_edit") else "PREGUNTA"
        return value.upper() if value else "PREGUNTA"

    def _answer_label(self) -> str:
        value = self.answer_label_edit.text().strip() if hasattr(self, "answer_label_edit") else "RESPUESTA"
        return value.upper() if value else "RESPUESTA"

    def _label_for_role(self, role: str) -> str:
        return self._question_label() if role == "question" else self._answer_label()

    def _sync_pr_button_names(self) -> None:
        if hasattr(self, "mark_question_button"):
            self.mark_question_button.setText(f"\U0001F3F7\ufe0f {self._question_label()}")
        if hasattr(self, "mark_answer_button"):
            self.mark_answer_button.setText(f"\U0001F3F7\ufe0f {self._answer_label()}")

    def _apply_custom_labels_to_pr_blocks(self, update_text: bool = True) -> None:
        for block in self.pr_blocks:
            role = str(block.get("role") or role_from_label(str(block.get("label", ""))))
            if role not in {"question", "answer"}:
                role = "question"
            block["role"] = role
            block["label"] = self._label_for_role(role)
        self._sync_pr_button_names()
        if update_text and self.pr_blocks:
            row = self.segments_list.currentRow() if hasattr(self, "segments_list") else None
            self._refresh_pr_list(keep_row=row)
            self._update_text_from_pr_blocks()

    def _pr_label_names_changed(self) -> None:
        self._sync_pr_button_names()
        if self.pr_blocks:
            self._apply_custom_labels_to_pr_blocks(update_text=True)

    def _make_pr_blocks_from_result(self) -> None:
        self.pr_blocks = []
        self.pr_notes = []
        if not self.last_result:
            return
        text = str(self.last_result.get("text", "")).strip()
        segments = self.last_result.get("segments", []) or []
        cleanup = bool(self.settings.get("clean_text", True))
        mode = FORMAT_MODES.get(str(self.settings.get("format", "Texto normal")), "normal")
        if mode == "qr_auto":
            self.pr_blocks, self.pr_notes = build_qr_blocks(segments, text, cleanup=cleanup)
        elif mode == "qr_call":
            self.pr_blocks, self.pr_notes = build_call_pr_blocks(segments, text, cleanup=cleanup)
        self._apply_custom_labels_to_pr_blocks(update_text=False)

    def _show_raw_segments(self, segments: List[Dict[str, Any]]) -> None:
        self.segments_list.setUpdatesEnabled(False)
        self.segments_list.clear()
        total = len(segments) if isinstance(segments, list) else 0
        for seg in (segments or [])[:MAX_LIST_ITEMS_RENDER]:
            if not isinstance(seg, dict):
                continue
            start = seconds_to_hhmmss(seg.get("start", 0))
            end = seconds_to_hhmmss(seg.get("end", 0))
            segment_text = str(seg.get("text", "")).strip()
            self.segments_list.addItem(f"[{start} - {end}] {segment_text}")
        if total > MAX_LIST_ITEMS_RENDER:
            self.segments_list.addItem(f"… {total - MAX_LIST_ITEMS_RENDER} segmentos más no se muestran en la lista para evitar cuelgues. Están guardados y se exportan igual.")
        self.segments_list.setUpdatesEnabled(True)
        self._set_pr_edit_enabled(False)

    def _format_block_item(self, idx: int, block: Dict[str, Any]) -> str:
        label = str(block.get("label", "PREGUNTA")).upper()
        start = seconds_to_hhmmss(block.get("start", 0))
        end = seconds_to_hhmmss(block.get("end", 0))
        text = re.sub(r"\s+", " ", str(block.get("text", "")).strip())
        if len(text) > 120:
            text = text[:117] + "..."
        return f"{idx + 1:02d}. {label}: [{start} - {end}] {text}"

    def _refresh_pr_list(self, keep_row: Optional[int] = None) -> None:
        self.segments_list.setUpdatesEnabled(False)
        self.segments_list.clear()
        total = len(self.pr_blocks) if isinstance(self.pr_blocks, list) else 0
        for idx, block in enumerate((self.pr_blocks or [])[:MAX_LIST_ITEMS_RENDER]):
            if isinstance(block, dict):
                self.segments_list.addItem(self._format_block_item(idx, block))
        if total > MAX_LIST_ITEMS_RENDER:
            self.segments_list.addItem(f"… {total - MAX_LIST_ITEMS_RENDER} bloques más no se muestran en la lista para evitar cuelgues. Están guardados y se exportan igual.")
        if self.pr_blocks and self.segments_list.count():
            row = keep_row if keep_row is not None else 0
            row = max(0, min(row, min(total, MAX_LIST_ITEMS_RENDER) - 1))
            self.segments_list.setCurrentRow(row)
        self.segments_list.setUpdatesEnabled(True)
        self._set_pr_edit_enabled(bool(self.pr_blocks))

    def _update_text_from_pr_blocks(self) -> None:
        if not self.pr_blocks:
            return
        self.text_edit.setPlainText(format_pr_blocks(self.pr_blocks, self.pr_notes, include_auto_note=self._current_pr_note_flag()))

    def _set_pr_edit_enabled(self, enabled: bool) -> None:
        has_row = enabled and self.segments_list.currentRow() >= 0
        for name in ["edit_pr_block_top_button", "edit_pr_block_button", "speakers_button", "mark_question_button", "mark_answer_button", "toggle_pr_button", "regen_pr_button"]:
            button = getattr(self, name, None)
            if button is not None:
                if name in {"edit_pr_block_top_button", "edit_pr_block_button", "speakers_button"}:
                    # Siempre queda visible/clickeable: si no hay bloque, el modal avisa qué falta.
                    button.setEnabled(True)
                else:
                    button.setEnabled(enabled if name == "regen_pr_button" else has_row)

    def _pr_selection_changed(self, _row: int) -> None:
        self._set_pr_edit_enabled(bool(self.pr_blocks))

    def set_selected_pr_role(self, role: str) -> None:
        if not self.pr_blocks:
            QMessageBox.information(self, "Sin P/R editable", "Primero aplica un formato P/R sobre la transcripcion.")
            return
        row = self.segments_list.currentRow()
        if row < 0 or row >= len(self.pr_blocks):
            QMessageBox.information(self, "Sin bloque seleccionado", "Selecciona un bloque de la lista para modificarlo.")
            return
        role = "question" if role == "question" else "answer"
        self.pr_blocks[row]["role"] = role
        self.pr_blocks[row]["label"] = self._label_for_role(role)
        self._refresh_pr_list(keep_row=row)
        self._update_text_from_pr_blocks()
        self.status_label.setText(f"Bloque {row + 1} marcado como {self.pr_blocks[row]['label']}.")
        self._autosave_transcription("bloque_marcado")

    def set_selected_pr_label(self, label: str) -> None:
        # Compatibilidad interna con versiones anteriores.
        self.set_selected_pr_role(role_from_label(label))

    def toggle_selected_pr_label(self) -> None:
        if not self.pr_blocks:
            return
        row = self.segments_list.currentRow()
        if row < 0 or row >= len(self.pr_blocks):
            return
        current_role = str(self.pr_blocks[row].get("role") or role_from_label(str(self.pr_blocks[row].get("label", "PREGUNTA"))))
        new_role = "answer" if current_role == "question" else "question"
        self.pr_blocks[row]["role"] = new_role
        self.pr_blocks[row]["label"] = self._label_for_role(new_role)
        self._refresh_pr_list(keep_row=row)
        self._update_text_from_pr_blocks()
        self.status_label.setText(f"Bloque {row + 1} alternado a {self.pr_blocks[row]['label']}.")
        self._autosave_transcription("bloque_alternado")

    def openEditBlockModal(self, blockId: Optional[int] = None) -> None:
        if isinstance(blockId, bool):
            blockId = None
        row = self.segments_list.currentRow() if blockId is None else int(blockId)
        print("CLICK EN EDITAR BLOQUE", row)
        self.open_pr_block_editor(row)

    def open_pr_block_editor(self, blockId: Optional[int] = None) -> None:
        row = self.segments_list.currentRow() if blockId is None else int(blockId)
        print("ABRIENDO MODAL EDITAR BLOQUE", row)
        dialog = QDialog(self)
        dialog.setObjectName("prEditDialog")
        dialog.setWindowTitle("Editar bloque de transcripción")
        dialog.setModal(True)
        dialog.resize(900, 680)

        layout = QVBoxLayout(dialog)
        layout.setContentsMargins(18, 18, 18, 18)
        layout.setSpacing(12)

        title = QLabel("Editar bloque de transcripción")
        title.setObjectName("sectionTitle")
        layout.addWidget(title)

        if not self.pr_blocks or row < 0 or row >= len(self.pr_blocks):
            message = QLabel("Seleccioná un bloque para editar.")
            message.setObjectName("warning")
            message.setWordWrap(True)
            layout.addWidget(message)
            actions_frame = QFrame()
            actions_frame.setObjectName("modalActions")
            actions = FlowLayout(actions_frame, margin=0, h_spacing=8, v_spacing=8)
            cancel_button = QPushButton("Cancelar")
            cancel_button.setObjectName("cancelButton")
            cancel_button.setMinimumHeight(38)
            cancel_button.clicked.connect(lambda: self.closeEditBlockModal(dialog))
            actions.addWidget(cancel_button)
            layout.addWidget(actions_frame)
            self._center_dialog(dialog)
            dialog.exec()
            return

        state = {"row": row}
        block = self.pr_blocks[row]
        block.setdefault("_original_text", str(block.get("text", "")))

        current = QLabel("")
        current.setObjectName("miniHelp")
        layout.addWidget(current)

        block_label = QLabel("Bloque actual")
        block_label.setObjectName("controlLabel")
        layout.addWidget(block_label)

        block_combo = QComboBox()
        for idx, item in enumerate(self.pr_blocks):
            block_combo.addItem(self._format_block_item(idx, item), idx)
        block_combo.setCurrentIndex(row)
        layout.addWidget(block_combo)

        player_frame = QFrame()
        player_frame.setObjectName("audioFragmentBox")
        player_layout = QHBoxLayout(player_frame)
        player_layout.setContentsMargins(10, 8, 10, 8)
        player_layout.setSpacing(8)

        fragment_label = QLabel("Fragmento: --:--:-- - --:--:--")
        fragment_label.setObjectName("miniHelp")
        player_layout.addWidget(fragment_label, stretch=1)

        play_button = QPushButton("▶ Escuchar fragmento")
        play_button.setObjectName("secondaryButton")
        play_button.setMinimumHeight(34)
        player_layout.addWidget(play_button)

        stop_button = QPushButton("⏹ Detener")
        stop_button.setObjectName("cancelButton")
        stop_button.setMinimumHeight(34)
        player_layout.addWidget(stop_button)

        speed_combo = QComboBox()
        speed_combo.setObjectName("smallInput")
        speed_combo.setToolTip("Velocidad de escucha del fragmento")
        speed_combo.addItems(["1x", "1.25x", "1.5x", "2x"])
        speed_combo.setCurrentText("1x")
        speed_combo.setMaximumWidth(90)
        player_layout.addWidget(speed_combo)

        seek_word_button = QPushButton("↪ Ir a palabra seleccionada")
        seek_word_button.setObjectName("secondaryButton")
        seek_word_button.setMinimumHeight(34)
        seek_word_button.setToolTip("Con timestamps por palabra: ubicá el cursor sobre una palabra y salta el audio a ese punto.")
        player_layout.addWidget(seek_word_button)

        layout.addWidget(player_frame)

        player_state: Dict[str, Any] = {"player": None, "timer": None, "end_ms": 0}

        def update_fragment_label(row_value: int) -> None:
            if 0 <= row_value < len(self.pr_blocks):
                selected = self.pr_blocks[row_value]
                start_s = float(selected.get("start", 0) or 0)
                end_s = float(selected.get("end", start_s) or start_s)
                fragment_label.setText(f"Fragmento: {seconds_to_hhmmss(start_s)} - {seconds_to_hhmmss(end_s)}")

        if QMediaPlayer is None or QAudioOutput is None:
            play_button.setEnabled(False)
            stop_button.setEnabled(False)
            seek_word_button.setEnabled(False)
            fragment_label.setText("Reproductor no disponible: falta QtMultimedia en PySide6.")
        elif not self.selected_file or not Path(self.selected_file).exists():
            play_button.setEnabled(False)
            stop_button.setEnabled(False)
            seek_word_button.setEnabled(False)
            fragment_label.setText("Reproductor no disponible: primero cargá el audio/video original.")
        else:
            player = QMediaPlayer(dialog)
            audio_output = QAudioOutput(dialog)
            audio_output.setVolume(0.9)
            player.setAudioOutput(audio_output)
            player.setSource(QUrl.fromLocalFile(self.selected_file))
            timer = QTimer(dialog)
            timer.setInterval(120)
            player_state.update({"player": player, "timer": timer})

            def stop_fragment() -> None:
                timer.stop()
                player.stop()
                try:
                    clear_audio_highlight()
                except Exception:
                    pass

            def play_fragment() -> None:
                current_row = int(state.get("row", row))
                if current_row < 0 or current_row >= len(self.pr_blocks):
                    QMessageBox.information(dialog, "Sin bloque", "Seleccioná un bloque para escuchar el fragmento.")
                    return
                selected = self.pr_blocks[current_row]
                start_s = max(float(selected.get("start", 0) or 0), 0.0)
                end_s = max(float(selected.get("end", start_s) or start_s), start_s)
                if end_s <= start_s:
                    end_s = start_s + 3.0
                player_state["end_ms"] = int(end_s * 1000)
                try:
                    player.setPlaybackRate(float(speed_combo.currentText().replace("x", "")))
                except Exception:
                    pass
                player.setPosition(int(start_s * 1000))
                player.play()
                timer.start()

            def auto_stop_fragment() -> None:
                if player.position() >= int(player_state.get("end_ms", 0)):
                    stop_fragment()

            play_button.clicked.connect(play_fragment)
            stop_button.clicked.connect(stop_fragment)
            # La conexión del timer se completa después de crear el editor de texto,
            # porque también actualiza el resaltado de la palabra/sección actual.
            dialog.finished.connect(lambda _result: stop_fragment())

        label_text = QLabel("Etiqueta / interlocutor")
        label_text.setObjectName("controlLabel")
        layout.addWidget(label_text)

        label_combo = QComboBox()
        label_combo.setEditable(True)
        label_options = [
            self._question_label(),
            self._answer_label(),
            "OPERADOR",
            "LLAMANTE",
            "ENTREVISTADOR",
            "ENTREVISTADO",
        ]
        label_combo.addItems(label_options)
        layout.addWidget(label_combo)

        text_label = QLabel("Texto del bloque")
        text_label.setObjectName("controlLabel")
        layout.addWidget(text_label)

        text_edit = QTextEdit()
        text_edit.setObjectName("modalTextEdit")
        text_edit.setMinimumHeight(175)
        layout.addWidget(text_edit, stretch=1)

        follow_label = QLabel("Seguimiento: si hay timestamps por palabra, se resalta la palabra exacta; si no, usa aproximación por sección.")
        follow_label.setObjectName("miniHelp")
        follow_label.setWordWrap(True)
        layout.addWidget(follow_label)

        def clear_audio_highlight() -> None:
            try:
                text_edit.setExtraSelections([])
                follow_label.setText("Seguimiento: si hay timestamps por palabra, se resalta la palabra exacta; si no, usa aproximación por sección.")
            except Exception:
                pass

        def update_audio_highlight() -> None:
            player_obj = player_state.get("player")
            if player_obj is None:
                return
            current_row = int(state.get("row", row))
            if current_row < 0 or current_row >= len(self.pr_blocks):
                clear_audio_highlight()
                return
            selected = self.pr_blocks[current_row]
            start_s = max(float(selected.get("start", 0) or 0), 0.0)
            end_s = max(float(selected.get("end", start_s) or start_s), start_s + 0.01)
            pos_s = float(player_obj.position()) / 1000.0
            if pos_s < start_s:
                clear_audio_highlight()
                return
            current_text = text_edit.toPlainText()
            matches = list(re.finditer(r"\S+", current_text))
            if not matches:
                clear_audio_highlight()
                return

            timed_words = normalize_word_items(selected.get("words", []))
            timed_index = words_for_time(timed_words, pos_s) if timed_words else None
            exact_mode = timed_index is not None

            if exact_mode:
                word_index = min(len(matches) - 1, max(0, int(timed_index)))
                section_size = 1
            else:
                progress = max(0.0, min(1.0, (pos_s - start_s) / max(end_s - start_s, 0.01)))
                word_index = min(len(matches) - 1, max(0, int(progress * len(matches))))
                section_size = max(1, min(5, len(matches) // 18 + 1))

            start_index = max(0, min(word_index, len(matches) - 1))
            end_index = min(len(matches) - 1, start_index + section_size - 1)
            cursor = QTextCursor(text_edit.document())
            cursor.setPosition(matches[start_index].start())
            cursor.setPosition(matches[end_index].end(), QTextCursor.KeepAnchor)
            fmt = QTextCharFormat()
            fmt.setBackground(QColor(251, 146, 60, 145 if exact_mode else 105))
            fmt.setForeground(QColor(255, 255, 255))
            selection = QTextEdit.ExtraSelection()
            selection.cursor = cursor
            selection.format = fmt
            text_edit.setExtraSelections([selection])
            preview = current_text[matches[start_index].start():matches[end_index].end()]
            mode_text = "exacto" if exact_mode else "aprox."
            follow_label.setText(f"Seguimiento {mode_text}: {seconds_to_hhmmss(pos_s)} → «{preview}»")

        def seek_to_selected_word() -> None:
            player_obj = player_state.get("player")
            if player_obj is None:
                return
            current_row = int(state.get("row", row))
            if current_row < 0 or current_row >= len(self.pr_blocks):
                return
            selected = self.pr_blocks[current_row]
            timed_words = normalize_word_items(selected.get("words", []))
            if not timed_words:
                QMessageBox.information(dialog, "Sin palabras sincronizadas", "Este bloque no tiene timestamps por palabra. Reintentá transcribir con Faster Whisper actualizado.")
                return
            current_text = text_edit.toPlainText()
            matches = list(re.finditer(r"\S+", current_text))
            if not matches:
                return
            cursor_pos = text_edit.textCursor().position()
            selected_index = 0
            for idx_m, match in enumerate(matches):
                if match.start() <= cursor_pos <= match.end():
                    selected_index = idx_m
                    break
                if cursor_pos > match.end():
                    selected_index = idx_m
            selected_index = min(selected_index, len(timed_words) - 1)
            target_ms = int(float(timed_words[selected_index].get("start", selected.get("start", 0)) or 0) * 1000)
            player_obj.setPosition(max(0, target_ms))
            update_audio_highlight()

        seek_word_button.clicked.connect(seek_to_selected_word)

        timer_obj = player_state.get("timer")
        player_obj = player_state.get("player")
        if timer_obj is not None and player_obj is not None:
            def on_audio_timer() -> None:
                update_audio_highlight()
                if player_obj.position() >= int(player_state.get("end_ms", 0)):
                    timer_obj.stop()
                    player_obj.stop()
                    clear_audio_highlight()
            timer_obj.timeout.connect(on_audio_timer)

        other_blocks_label = QLabel("Otras frases / bloques de la transcripción")
        other_blocks_label.setObjectName("controlLabel")
        other_blocks_label.setToolTip("Tocá una frase de la lista para cargarla arriba y modificar etiqueta o texto.")
        layout.addWidget(other_blocks_label)

        modal_blocks_list = QListWidget()
        modal_blocks_list.setObjectName("modalBlockList")
        modal_blocks_list.setMinimumHeight(105)
        modal_blocks_list.setMaximumHeight(160)
        modal_blocks_list.setToolTip("Tocá una frase para seleccionarla y editarla en este mismo modal")
        for idx, item in enumerate(self.pr_blocks):
            modal_blocks_list.addItem(self._format_block_item(idx, item))
        modal_blocks_list.setCurrentRow(row)
        layout.addWidget(modal_blocks_list)

        def load_block(index: int) -> None:
            row_value = block_combo.itemData(index)
            if row_value is None:
                return
            state["row"] = int(row_value)
            selected = self.pr_blocks[state["row"]]
            selected.setdefault("_original_text", str(selected.get("text", "")))
            current.setText(f"Bloque {state['row'] + 1} de {len(self.pr_blocks)}")
            update_fragment_label(state["row"])
            player_obj = player_state.get("player")
            timer_obj = player_state.get("timer")
            if player_obj is not None:
                if timer_obj is not None:
                    timer_obj.stop()
                player_obj.stop()
            clear_audio_highlight()
            current_label = str(selected.get("label", self._label_for_role(str(selected.get("role", "question"))))).upper()
            if current_label and label_combo.findText(current_label) < 0:
                label_combo.addItem(current_label)
            label_combo.setCurrentText(current_label)
            text_edit.setPlainText(str(selected.get("text", "")))
            word_total = len(normalize_word_items(selected.get("words", [])))
            if word_total:
                follow_label.setText(f"Seguimiento exacto disponible: {word_total} palabra/s sincronizada/s.")
            else:
                follow_label.setText("Seguimiento aproximado: este bloque no tiene timestamps por palabra.")
            modal_blocks_list.blockSignals(True)
            modal_blocks_list.setCurrentRow(state["row"])
            modal_blocks_list.blockSignals(False)

        def select_block_from_bottom_list(item) -> None:
            selected_row = modal_blocks_list.row(item)
            if 0 <= selected_row < block_combo.count():
                block_combo.setCurrentIndex(selected_row)

        modal_blocks_list.itemClicked.connect(select_block_from_bottom_list)
        modal_blocks_list.itemDoubleClicked.connect(select_block_from_bottom_list)
        block_combo.currentIndexChanged.connect(load_block)
        load_block(block_combo.currentIndex())

        actions_frame = QFrame()
        actions_frame.setObjectName("modalActions")
        actions = FlowLayout(actions_frame, margin=0, h_spacing=8, v_spacing=8)

        restore_button = QPushButton("Restaurar texto original")
        restore_button.setObjectName("secondaryButton")
        restore_button.setMinimumHeight(38)
        restore_button.clicked.connect(lambda: self.restoreOriginalBlock(state["row"], text_edit))
        actions.addWidget(restore_button)

        cancel_button = QPushButton("Cancelar")
        cancel_button.setObjectName("cancelButton")
        cancel_button.setMinimumHeight(38)
        cancel_button.clicked.connect(lambda: self.closeEditBlockModal(dialog))
        actions.addWidget(cancel_button)

        save_button = QPushButton("Guardar cambios")
        save_button.setObjectName("primaryButton")
        save_button.setMinimumHeight(38)
        save_button.clicked.connect(lambda: self.saveEditedBlock(dialog, state["row"], label_combo.currentText(), text_edit.toPlainText()))
        actions.addWidget(save_button)

        layout.addWidget(actions_frame)
        self._center_dialog(dialog)
        dialog.exec()

    def closeEditBlockModal(self, dialog: QDialog) -> None:
        dialog.reject()

    def restoreOriginalBlock(self, row: int, text_edit: QTextEdit) -> None:
        if 0 <= row < len(self.pr_blocks):
            text_edit.setPlainText(str(self.pr_blocks[row].get("_original_text", self.pr_blocks[row].get("text", ""))))

    def saveEditedBlock(self, dialog: QDialog, row: int, label: str, text: str) -> None:
        self.save_pr_block_editor(dialog, row, label, text)

    def save_pr_block_editor(self, dialog: QDialog, row: int, label: str, text: str) -> None:
        if row < 0 or row >= len(self.pr_blocks):
            dialog.reject()
            return
        label = (label or "").strip().upper() or self._question_label()
        role = "question" if label == self._question_label() else "answer" if label == self._answer_label() else role_from_label(label)
        self.pr_blocks[row]["role"] = role
        self.pr_blocks[row]["label"] = label
        self.pr_blocks[row]["text"] = text.strip()
        self._refresh_pr_list(keep_row=row)
        self._update_text_from_pr_blocks()
        self.status_label.setText(f"Bloque {row + 1} editado.")
        self._autosave_transcription("bloque_editado")
        dialog.accept()


    def _speaker_id_for_block_index(self, index: int, block: Dict[str, Any]) -> str:
        existing = str(block.get("speaker_id") or "").strip()
        if existing in {"SPEAKER_1", "SPEAKER_2"}:
            return existing
        role = str(block.get("role") or "").lower()
        if role == "question":
            return "SPEAKER_1"
        if role == "answer":
            return "SPEAKER_2"
        return "SPEAKER_1" if index % 2 == 0 else "SPEAKER_2"

    def _speaker_label(self, speaker_id: str) -> str:
        return str(self.speaker_names.get(speaker_id, speaker_id)).strip() or speaker_id

    def _apply_speaker_names_to_blocks(self, update_text: bool = True) -> None:
        """Aplica nombres generales Speaker 1/Speaker 2 sin perder textos ni timestamps."""
        for idx, block in enumerate(self.pr_blocks):
            speaker_id = self._speaker_id_for_block_index(idx, block)
            block["speaker_id"] = speaker_id
            block["label"] = self._speaker_label(speaker_id)
            block["role"] = "question" if speaker_id == "SPEAKER_1" else "answer"
        if self.pr_blocks:
            self.question_label_edit.blockSignals(True)
            self.answer_label_edit.blockSignals(True)
            self.question_label_edit.setText(self._speaker_label("SPEAKER_1"))
            self.answer_label_edit.setText(self._speaker_label("SPEAKER_2"))
            self.question_label_edit.blockSignals(False)
            self.answer_label_edit.blockSignals(False)
            self._refresh_pr_list(keep_row=max(0, self.segments_list.currentRow()))
            if update_text:
                self._update_text_from_pr_blocks()

    def _build_speaker_blocks_from_segments(self) -> None:
        """Diarización básica general: alterna interlocutor por pausas/segmentos como base editable.

        No pretende identificar voces reales; deja una propuesta rápida Speaker 1/Speaker 2 para corregir.
        """
        if not self.last_result:
            QMessageBox.information(self, "Sin transcripción", "Primero realizá una transcripción.")
            return
        segments = self.last_result.get("segments", []) or []
        text = str(self.last_result.get("text", "")).strip()
        if not segments:
            if text:
                self.pr_blocks = [{
                    "role": "question",
                    "speaker_id": "SPEAKER_1",
                    "label": self._speaker_label("SPEAKER_1"),
                    "text": text,
                    "start": 0,
                    "end": 0,
                    "words": [],
                }]
                self.pr_notes = ["Diarización básica creada desde texto completo, sin segmentos temporales."]
            else:
                QMessageBox.information(self, "Sin segmentos", "No hay segmentos para generar interlocutores.")
                return
        else:
            self.pr_blocks = []
            current_speaker = "SPEAKER_1"
            previous_end = None
            for idx, seg in enumerate(segments):
                seg_text = re.sub(r"\s+", " ", str(seg.get("text", "")).strip())
                if not seg_text:
                    continue
                start = float(seg.get("start", 0) or 0)
                end = float(seg.get("end", start) or start)
                pause = (start - previous_end) if previous_end is not None else 0
                # Heurística general: alterna en pausas marcadas o segmentos con signo de pregunta.
                if idx > 0 and (pause >= 1.2 or str(self.pr_blocks[-1].get("text", "")).strip().endswith("?")):
                    current_speaker = "SPEAKER_2" if current_speaker == "SPEAKER_1" else "SPEAKER_1"
                self.pr_blocks.append({
                    "role": "question" if current_speaker == "SPEAKER_1" else "answer",
                    "speaker_id": current_speaker,
                    "label": self._speaker_label(current_speaker),
                    "text": seg_text,
                    "start": start,
                    "end": end,
                    "words": normalize_word_items(seg.get("words", [])),
                })
                previous_end = end
            self.pr_notes = ["Interlocutores generados automáticamente de forma básica. Revisar y corregir si hace falta."]
        self.settings["format"] = "P/R alternado simple"
        self._apply_speaker_names_to_blocks(update_text=True)
        self._update_pr_format_button()
        self.status_label.setText("Interlocutores básicos generados. Podés renombrar Speaker 1 / Speaker 2 y editar cada bloque.")
        self._autosave_transcription("interlocutores_generados")

    def open_speakers_modal(self) -> None:
        dialog = QDialog(self)
        dialog.setObjectName("speakersDialog")
        dialog.setWindowTitle("Interlocutores")
        dialog.setModal(True)
        dialog.resize(560, 360)

        layout = QVBoxLayout(dialog)
        layout.setContentsMargins(18, 18, 18, 18)
        layout.setSpacing(12)

        title = QLabel("👥 Interlocutores")
        title.setObjectName("sectionTitle")
        layout.addWidget(title)

        help_text = QLabel("Función general: genera o renombra Speaker 1 / Speaker 2. Es una diarización básica editable, no identificación biométrica de voces.")
        help_text.setObjectName("miniHelp")
        help_text.setWordWrap(True)
        layout.addWidget(help_text)

        grid_frame = QFrame()
        grid = QGridLayout(grid_frame)
        grid.setContentsMargins(0, 0, 0, 0)
        grid.setHorizontalSpacing(10)
        grid.setVerticalSpacing(10)

        speaker1_edit = QLineEdit(self._speaker_label("SPEAKER_1"))
        speaker1_edit.setObjectName("smallInput")
        speaker1_edit.setPlaceholderText("Ej.: Speaker 1, Entrevistador, Host")
        speaker2_edit = QLineEdit(self._speaker_label("SPEAKER_2"))
        speaker2_edit.setObjectName("smallInput")
        speaker2_edit.setPlaceholderText("Ej.: Speaker 2, Invitado, Persona 2")

        label1 = QLabel("Speaker 1")
        label1.setObjectName("controlLabel")
        label2 = QLabel("Speaker 2")
        label2.setObjectName("controlLabel")
        grid.addWidget(label1, 0, 0)
        grid.addWidget(speaker1_edit, 0, 1)
        grid.addWidget(label2, 1, 0)
        grid.addWidget(speaker2_edit, 1, 1)
        layout.addWidget(grid_frame)

        stats = QLabel(f"Bloques actuales: {len(self.pr_blocks)}")
        stats.setObjectName("progressMeta")
        layout.addWidget(stats)

        actions_frame = QFrame()
        actions_frame.setObjectName("modalActions")
        actions = FlowLayout(actions_frame, margin=0, h_spacing=8, v_spacing=8)

        detect_button = QPushButton("Detectar/crear bloques")
        detect_button.setObjectName("secondaryButton")
        detect_button.setMinimumHeight(38)
        detect_button.setToolTip("Crea una base Speaker 1 / Speaker 2 desde los segmentos actuales")
        detect_button.clicked.connect(lambda: (self._save_speaker_names_from_edits(speaker1_edit, speaker2_edit), self._build_speaker_blocks_from_segments(), stats.setText(f"Bloques actuales: {len(self.pr_blocks)}")))
        actions.addWidget(detect_button)

        apply_button = QPushButton("Aplicar nombres")
        apply_button.setObjectName("primaryButton")
        apply_button.setMinimumHeight(38)
        apply_button.clicked.connect(lambda: (self._save_speaker_names_from_edits(speaker1_edit, speaker2_edit), self._apply_speaker_names_to_blocks(update_text=True), self._autosave_transcription("interlocutores_renombrados"), dialog.accept()))
        actions.addWidget(apply_button)

        cancel_button = QPushButton("Cancelar")
        cancel_button.setObjectName("cancelButton")
        cancel_button.setMinimumHeight(38)
        cancel_button.clicked.connect(dialog.reject)
        actions.addWidget(cancel_button)
        layout.addWidget(actions_frame)

        self._center_dialog(dialog)
        dialog.exec()

    def _save_speaker_names_from_edits(self, speaker1_edit: QLineEdit, speaker2_edit: QLineEdit) -> None:
        self.speaker_names["SPEAKER_1"] = speaker1_edit.text().strip() or "Speaker 1"
        self.speaker_names["SPEAKER_2"] = speaker2_edit.text().strip() or "Speaker 2"

    def marcarEtiqueta1(self) -> None:
        self.set_selected_pr_role("question")

    def marcarEtiqueta2(self) -> None:
        self.set_selected_pr_role("answer")

    def alternarEtiquetas(self) -> None:
        self.toggle_selected_pr_label()

    def regenerarPR(self) -> None:
        self.apply_qr_format()

    def abrirModalEditarBloque(self, _bloque: Optional[Dict[str, Any]] = None) -> None:
        self.open_pr_block_editor()

    def select_file(self) -> None:
        filter_text = "Audio/Video (*.mp4 *.avi *.mkv *.mov *.webm *.mp3 *.wav *.m4a *.aac *.flac *.ogg *.wma);;Todos los archivos (*.*)"
        file_path, _ = QFileDialog.getOpenFileName(self, "Seleccionar audio o video", str(Path.home()), filter_text)
        if file_path:
            self.set_file(file_path)

    def set_file(self, file_path: str) -> None:
        path = Path(file_path)
        if not path.exists():
            QMessageBox.warning(self, "Archivo no encontrado", "El archivo seleccionado no existe.")
            return
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            QMessageBox.warning(self, "Formato no soportado", f"Formato no soportado: {path.suffix}")
            return
        self.selected_file = str(path)
        self.file_label.setText(str(path))
        self.file_label.setToolTip(str(path))
        self._detect_selected_file_duration()
        duration_text = seconds_to_duration_label(self.media_duration_seconds)
        self.status_label.setText(f"Archivo cargado. Listo para transcribir. Duración: {duration_text}")
        self._set_progress(0)

    def _resolve_language(self) -> Optional[str]:
        selected = str(self.settings.get("language", "Espanol"))
        value = LANGUAGES.get(selected)
        if value == "manual":
            manual = str(self.settings.get("manual_language", "")).strip()
            return manual or None
        return value

    def start_transcription(self) -> None:
        if not self.selected_file:
            QMessageBox.information(self, "Falta archivo", "Primero selecciona un audio o video.")
            return
        if shutil.which("ffmpeg") is None:
            QMessageBox.warning(self, "FFmpeg requerido", "Instala FFmpeg y asegurate de que este en PATH antes de transcribir.")
            return

        model = str(self.settings.get("model", "small"))
        task = TASKS.get(str(self.settings.get("task", "Transcribir")), "transcribe")
        language = self._resolve_language()
        profile = PERFORMANCE_PROFILES.get(str(self.settings.get("performance", "Rapido")), "fast")
        device_mode = DEVICE_OPTIONS.get(str(self.settings.get("device", "Automatico")), "auto")
        engine = ENGINE_OPTIONS.get(str(self.settings.get("engine", "Faster Whisper (mas rapido)")), "faster")
        compute_type = COMPUTE_OPTIONS.get(str(self.settings.get("compute", "Auto")), "auto")
        optimize_audio = bool(self.settings.get("optimize_audio", True))
        thread_text = str(self.settings.get("threads", "Auto"))
        cpu_threads = int(thread_text) if thread_text.isdigit() else 0
        chunk_seconds = CHUNK_OPTIONS.get(str(self.settings.get("chunk", "10 min")), 0)

        self.last_result = None
        self.pr_blocks = []
        self.pr_notes = []
        self.text_edit.clear()
        if hasattr(self, "search_input"):
            self.search_input.clear()
        self.segments_list.clear()
        self._set_export_enabled(False)
        self._set_busy(True)
        self._set_progress(0)
        self.status_label.setText("Iniciando...")

        self.worker = TranscriptionWorker(
            self.selected_file,
            model,
            language,
            task,
            profile,
            device_mode,
            engine,
            compute_type,
            optimize_audio,
            cpu_threads,
            chunk_seconds,
        )
        self.worker.status_changed.connect(self.status_label.setText)
        self.worker.progress_changed.connect(self._set_progress)
        self.worker.finished_ok.connect(self._transcription_finished)
        self.worker.failed.connect(self._transcription_failed)
        self.worker.finished.connect(lambda: self._set_busy(False))
        self.worker.start()

    def cancel_transcription(self) -> None:
        if self.worker and self.worker.isRunning():
            self.worker.requestInterruption()
            self.status_label.setText("Cancelacion solicitada. Si Whisper esta procesando, puede demorar en detenerse.")
            self.cancel_button.setEnabled(False)

    def _transcription_finished(self, result: Dict[str, Any]) -> None:
        self.last_result = result
        text = str(result.get("text", "")).strip()
        segments = result.get("segments", []) or []
        word_count = sum(len(seg.get("words", []) or []) for seg in segments if isinstance(seg, dict))
        if self.media_duration_seconds is None:
            self.media_duration_seconds = duration_from_segments(segments)
            self._refresh_progress_meta(100)
        mode = FORMAT_MODES.get(str(self.settings.get("format", "Texto normal")), "normal")
        cleanup = bool(self.settings.get("clean_text", True))

        self.pr_blocks = []
        self.pr_notes = []
        if mode in {"qr_auto", "qr_call"}:
            self._make_pr_blocks_from_result()
            self._update_text_from_pr_blocks()
            self._refresh_pr_list()
        else:
            self.text_edit.setPlainText(build_formatted_text(mode, segments, text, cleanup=cleanup))
            self._show_raw_segments(segments)

        self._set_export_enabled(True)
        self.apply_qr_button.setEnabled(True)
        self._update_pr_format_button()
        runtime = result.get("_dpi_runtime", {}) or {}
        device = str(runtime.get("device", "")).upper() or "AUTO"
        chunks = int(runtime.get("chunks", 1) or 1)
        threads = runtime.get("cpu_threads", "Auto")
        extra = f" Fragmentos: {chunks}." if chunks > 1 else ""
        if mode in {"qr_auto", "qr_call"}:
            extra += " P/R editable activado."
        if word_count:
            extra += f" Palabras sincronizadas: {word_count}."
        engine = str(runtime.get("engine", "openai")).upper()
        compute = str(runtime.get("compute_type", "")).upper()
        self.status_label.setText(f"Finalizado. Modelo cacheado para la proxima. Motor: {engine}. Dispositivo: {device}. Precision: {compute}. Hilos: {threads}.{extra}")
        self._autosave_transcription("finalizada")

    def _transcription_failed(self, message: str) -> None:
        self.status_label.setText("Error durante la transcripcion.")
        QMessageBox.critical(self, "Error", message)

    def _update_pr_format_button(self) -> None:
        if not hasattr(self, "apply_qr_button"):
            return
        mode = FORMAT_MODES.get(str(self.settings.get("format", "Texto normal")), "normal")
        pr_active = bool(self.pr_blocks) and mode in {"qr_auto", "qr_call"}
        if pr_active:
            self.apply_qr_button.setText("Quitar P/R")
            self.apply_qr_button.setToolTip("Vuelve la transcripcion a texto normal, sin formato P/R")
        else:
            self.apply_qr_button.setText("Usar P/R")
            self.apply_qr_button.setToolTip("Convierte la transcripcion a bloques PREGUNTA/RESPUESTA editables")

    def toggle_pr_format(self) -> None:
        if not self.last_result:
            QMessageBox.information(self, "Sin transcripcion", "Primero realiza una transcripcion.")
            return
        mode = FORMAT_MODES.get(str(self.settings.get("format", "Texto normal")), "normal")
        if self.pr_blocks and mode in {"qr_auto", "qr_call"}:
            self.remove_pr_format()
        else:
            self.apply_qr_format()

    def remove_pr_format(self) -> None:
        if not self.last_result:
            return
        self.settings["format"] = "Texto normal"
        text = str(self.last_result.get("text", "")).strip()
        segments = self.last_result.get("segments", []) or []
        cleanup = bool(self.settings.get("clean_text", True))
        self.pr_blocks = []
        self.pr_notes = []
        self.text_edit.setPlainText(build_formatted_text("normal", segments, text, cleanup=cleanup))
        self._show_raw_segments(segments)
        self._update_pr_format_button()
        self.status_label.setText("Formato P/R quitado. La transcripcion vuelve a texto normal, sin bloques P/R.")
        self._autosave_transcription("pr_quitado")

    def apply_qr_format(self) -> None:
        if not self.last_result:
            QMessageBox.information(self, "Sin transcripcion", "Primero realiza una transcripcion.")
            return
        mode = FORMAT_MODES.get(str(self.settings.get("format", "Texto normal")), "normal")
        if mode == "normal":
            self.settings["format"] = "P/R inteligente llamada"
        self._make_pr_blocks_from_result()
        self._update_text_from_pr_blocks()
        self._refresh_pr_list()
        self._update_pr_format_button()
        self.status_label.setText("Formato P/R aplicado. Ahora podes seleccionar bloques y marcar PREGUNTA o RESPUESTA sin editar todo a mano.")
        self._autosave_transcription("pr_aplicado")

    def _set_busy(self, busy: bool) -> None:
        self.transcribe_button.setEnabled(not busy)
        self.cancel_button.setEnabled(busy)
        self.select_button.setEnabled(not busy)
        self.clear_button.setEnabled(not busy)
        self.performance_button.setEnabled(not busy)
        self.apply_qr_button.setEnabled((not busy) and bool(self.last_result))
        self._update_pr_format_button()
        if busy:
            self._set_pr_edit_enabled(False)
        else:
            self._set_pr_edit_enabled(bool(self.pr_blocks))

    def _set_export_enabled(self, enabled: bool) -> None:
        for button in [
            self.export_txt_button,
            self.export_srt_button,
            self.export_vtt_button,
            self.export_json_button,
            self.export_docx_button,
        ]:
            button.setEnabled(enabled)

    def clear_all(self) -> None:
        self.selected_file = None
        self.last_result = None
        self.pr_blocks = []
        self.pr_notes = []
        self.speaker_names = {"SPEAKER_1": "Speaker 1", "SPEAKER_2": "Speaker 2"}
        self.current_backup_path = None
        self.media_duration_seconds = None
        self.file_label.setText("Ningun archivo seleccionado")
        self.file_label.setToolTip("Ruta del archivo seleccionado")
        self.text_edit.clear()
        if hasattr(self, "search_input"):
            self.search_input.clear()
        self.search_matches = []
        self.search_match_index = -1
        self._clear_search_highlights()
        self.segments_list.clear()
        self._set_progress(0)
        self.status_label.setText("Esperando archivo...")
        self.apply_qr_button.setEnabled(False)
        self._update_pr_format_button()
        self._set_pr_edit_enabled(False)
        self._set_export_enabled(False)

    def _ensure_autosave_path(self) -> Path:
        autosave_dir().mkdir(parents=True, exist_ok=True)
        if self.current_backup_path is not None:
            try:
                self.current_backup_path.parent.mkdir(parents=True, exist_ok=True)
                return self.current_backup_path
            except Exception:
                self.current_backup_path = None
        base_name = Path(self.selected_file).stem if self.selected_file else "transcripcion"
        slug = safe_file_slug(base_name)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        short_hash = hashlib.sha1(f"{self.selected_file or ''}|{stamp}".encode("utf-8", errors="ignore")).hexdigest()[:6]
        self.current_backup_path = autosave_dir() / f"{stamp}_{slug}_{short_hash}.artrans.json"
        return self.current_backup_path

    def _autosave_transcription(self, reason: str = "actualizada") -> None:
        """Guarda una copia silenciosa/importable sin interrumpir al operador."""
        if not self.last_result:
            return
        try:
            path = self._ensure_autosave_path()
            payload = {
                "schema": "AR_TRANSCRIPTOR_AUTOSAVE",
                "schema_version": AUTOSAVE_SCHEMA_VERSION,
                "app_name": APP_NAME,
                "app_version": APP_VERSION,
                "saved_at": datetime.now().isoformat(timespec="seconds"),
                "saved_reason": reason,
                "selected_file": self.selected_file or "",
                "file_exists": bool(self.selected_file and Path(self.selected_file).exists()),
                "media_duration_seconds": self.media_duration_seconds,
                "displayed_text": self.text_edit.toPlainText().strip() if hasattr(self, "text_edit") else "",
                "last_result": self.last_result,
                "pr_blocks": self.pr_blocks,
                "pr_notes": self.pr_notes,
                "speaker_names": self.speaker_names,
                "review_keywords": self.review_keywords,
                "settings": self.settings,
                "metadata": self._metadata(),
            }
            path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
            print(f"Copia AR_TRANSCRIPTOR guardada: {path}")
        except Exception as exc:
            # Copia discreta: no frena el trabajo ni molesta con ventanas emergentes.
            print(f"No se pudo guardar copia AR_TRANSCRIPTOR: {exc}")

    def _safe_set_transcription_text(self, text: str) -> None:
        """Carga texto grande sin disparar búsquedas/repaint pesados durante importación."""
        text = str(text or "")
        truncated = False
        if len(text) > MAX_IMPORTED_DISPLAY_CHARS:
            text = text[:MAX_IMPORTED_DISPLAY_CHARS] + "\n\n[Vista previa limitada para mantener fluida la app. La copia completa queda cargada internamente y puede exportarse.]"
            truncated = True
        old = self.text_edit.blockSignals(True) if hasattr(self, "text_edit") else False
        try:
            self.text_edit.setPlainText(text)
        finally:
            if hasattr(self, "text_edit"):
                self.text_edit.blockSignals(old)
        if truncated and hasattr(self, "status_label"):
            self.status_label.setText("Copia importada con vista previa limitada por tamaño.")

    def _normalize_imported_segments(self, segments: Any) -> List[Dict[str, Any]]:
        if not isinstance(segments, list):
            return []
        clean_segments: List[Dict[str, Any]] = []
        for seg in segments:
            if isinstance(seg, dict):
                clean_segments.append(seg)
        return clean_segments

    def import_saved_transcription(self) -> None:
        start_dir = str(autosave_dir()) if autosave_dir().exists() else str(Path.home())
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Importar copia de transcripción",
            start_dir,
            "Copias AR_TRANSCRIPTOR (*.artrans.json *.json);;JSON (*.json);;Todos (*.*)",
        )
        if not file_path:
            return
        path = Path(file_path)
        QApplication.setOverrideCursor(Qt.WaitCursor)
        self._bulk_loading_ui = True
        try:
            if not path.exists():
                raise FileNotFoundError("No se encontró el archivo seleccionado.")
            # Límite preventivo: evita que un JSON enorme deje la app sin responder.
            if path.stat().st_size > 80 * 1024 * 1024:
                raise ValueError("La copia pesa demasiado para importarse de forma segura desde la interfaz. Probá con una copia más liviana.")

            if hasattr(self, "status_label"):
                self.status_label.setText("Importando copia…")
            QApplication.processEvents()

            payload = json.loads(path.read_text(encoding="utf-8"))
            if not isinstance(payload, dict):
                raise ValueError("La copia no tiene formato JSON válido.")

            imported_settings = payload.get("settings")
            if isinstance(imported_settings, dict):
                self.settings.update(imported_settings)

            raw_result = payload.get("last_result") or payload.get("raw_result")
            if not isinstance(raw_result, dict):
                text = str(payload.get("text") or payload.get("displayed_text") or "").strip()
                segments = self._normalize_imported_segments(payload.get("segments"))
                raw_result = {"text": text, "segments": segments}
            else:
                raw_result["segments"] = self._normalize_imported_segments(raw_result.get("segments"))
                raw_result["text"] = str(raw_result.get("text", "") or "")

            self.last_result = raw_result
            self.selected_file = str(payload.get("selected_file") or raw_result.get("file_path") or "") or None
            saved_duration = payload.get("media_duration_seconds")
            try:
                self.media_duration_seconds = float(saved_duration) if saved_duration not in (None, "") else None
            except (TypeError, ValueError):
                self.media_duration_seconds = None
            if self.media_duration_seconds is None:
                self.media_duration_seconds = duration_from_segments(raw_result.get("segments") or [])

            self.current_backup_path = path
            if self.selected_file:
                self.file_label.setText(self.selected_file)
                self.file_label.setToolTip(self.selected_file)
            else:
                self.file_label.setText(f"Copia importada: {path.name}")
                self.file_label.setToolTip(str(path))

            imported_blocks = payload.get("pr_blocks")
            self.pr_blocks = [b for b in imported_blocks if isinstance(b, dict)] if isinstance(imported_blocks, list) else []
            imported_notes = payload.get("pr_notes")
            self.pr_notes = [str(n) for n in imported_notes] if isinstance(imported_notes, list) else []
            imported_speakers = payload.get("speaker_names")
            if isinstance(imported_speakers, dict):
                self.speaker_names.update({str(k): str(v) for k, v in imported_speakers.items() if str(k).startswith("SPEAKER_")})
            imported_keywords = payload.get("review_keywords")
            if isinstance(imported_keywords, list):
                self.review_keywords = [str(x) for x in imported_keywords if str(x).strip()]

            # Limpiar búsqueda antes de cargar texto: si quedó una búsqueda activa, PySide puede tardar muchísimo
            # resaltando miles de resultados en una importación grande.
            if hasattr(self, "search_input"):
                self.search_input.blockSignals(True)
                self.search_input.clear()
                self.search_input.blockSignals(False)
            self.search_matches = []
            self.search_match_index = -1
            self._clear_search_highlights()

            displayed_text = str(payload.get("displayed_text") or "").strip()
            if displayed_text:
                self._safe_set_transcription_text(displayed_text)
            elif self.pr_blocks:
                self._safe_set_transcription_text(format_pr_blocks(self.pr_blocks, self.pr_notes, include_auto_note=self._current_pr_note_flag()))
            else:
                text = str(self.last_result.get("text", "")).strip()
                segments = self.last_result.get("segments", []) or []
                cleanup = bool(self.settings.get("clean_text", True))
                self._safe_set_transcription_text(build_formatted_text("normal", segments, text, cleanup=cleanup))

            QApplication.processEvents()
            if self.pr_blocks:
                self._refresh_pr_list(keep_row=0)
            else:
                self._show_raw_segments(self.last_result.get("segments", []) or [])

            self._set_progress(100)
            self._set_export_enabled(True)
            self.apply_qr_button.setEnabled(True)
            self._update_pr_format_button()
            if hasattr(self, "progress_meta_label"):
                pct = "100%"
                dur = seconds_to_duration_label(self.media_duration_seconds) if self.media_duration_seconds else "--:--"
                self.progress_meta_label.setText(f"{pct} · Duración {dur}")
            self.status_label.setText(f"Copia importada: {path.name}")
        except Exception as exc:
            QMessageBox.critical(self, "No se pudo importar", str(exc))
        finally:
            self._bulk_loading_ui = False
            QApplication.restoreOverrideCursor()

    def _metadata(self) -> Dict[str, Any]:
        language = str(self.settings.get("language", "Espanol"))
        return {
            "archivo": self.selected_file or "",
            "duracion_audio_video": seconds_to_duration_label(self.media_duration_seconds),
            "duracion_segundos": self.media_duration_seconds,
            "modelo": self.settings.get("model", "small"),
            "idioma": language if language != "Otro/manual" else str(self.settings.get("manual_language", "")).strip(),
            "tarea": self.settings.get("task", "Transcribir"),
            "formato": self.settings.get("format", "Texto normal"),
            "perfil_rendimiento": self.settings.get("performance", "Rapido"),
            "dispositivo": self.settings.get("device", "Automatico"),
            "motor": self.settings.get("engine", "Faster Whisper (mas rapido)"),
            "precision_compute": self.settings.get("compute", "Auto"),
            "audio_optimizado": "Si" if bool(self.settings.get("optimize_audio", True)) else "No",
            "hilos_cpu_ffmpeg": self.settings.get("threads", "Auto"),
            "troceo_archivo": self.settings.get("chunk", "10 min"),
            "limpieza_basica": "Si" if bool(self.settings.get("clean_text", True)) else "No",
            "fecha_hora": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
        }

    def export_result(self, kind: str) -> None:
        if not self.last_result:
            QMessageBox.information(self, "Sin resultado", "Primero realiza una transcripcion.")
            return

        base = Path(self.selected_file).stem if self.selected_file else "transcripcion"
        suggested = f"{base}_transcripcion.{kind}"
        file_path, _ = QFileDialog.getSaveFileName(self, "Guardar exportacion", suggested, f"*.{kind}")
        if not file_path:
            return

        path = Path(file_path)
        if path.suffix.lower() != f".{kind}":
            path = path.with_suffix(f".{kind}")

        text = self.text_edit.toPlainText().strip() or str(self.last_result.get("text", "")).strip()
        segments = self.last_result.get("segments", []) or []
        metadata = self._metadata()

        try:
            if kind == "txt":
                path.write_text(build_txt(metadata, text, segments), encoding="utf-8")
            elif kind == "srt":
                path.write_text(build_srt(segments), encoding="utf-8")
            elif kind == "vtt":
                path.write_text(build_vtt(segments), encoding="utf-8")
            elif kind == "json":
                payload = {
                    "metadata": metadata,
                    "text": text,
                    "segments": segments,
                    "raw_result": self.last_result,
                    "pr_blocks": self.pr_blocks,
                    "pr_notes": self.pr_notes,
                    "speaker_names": self.speaker_names,
                }
                path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
            elif kind == "docx":
                self._export_docx(path, metadata, text, segments)
            else:
                raise ValueError("Tipo de exportacion no soportado.")

            QMessageBox.information(self, "Exportado", f"Archivo guardado en:\n{path}")
        except Exception as exc:
            QMessageBox.critical(self, "Error al exportar", str(exc))

    def _export_docx(self, path: Path, metadata: Dict[str, Any], text: str, segments: List[Dict[str, Any]]) -> None:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()
        styles = doc.styles
        styles["Normal"].font.name = "Arial"
        styles["Normal"].font.size = Pt(11)

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("TRANSCRIPCION DE AUDIO/VIDEO")
        run.bold = True
        run.font.size = Pt(14)

        doc.add_paragraph("")
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for key, value in [
            ("Archivo analizado", metadata.get("archivo", "")),
            ("Modelo utilizado", metadata.get("modelo", "")),
            ("Idioma", metadata.get("idioma", "")),
            ("Tarea", metadata.get("tarea", "")),
            ("Formato", metadata.get("formato", "")),
            ("Perfil rendimiento", metadata.get("perfil_rendimiento", "")),
            ("Dispositivo", metadata.get("dispositivo", "")),
            ("Motor", metadata.get("motor", "")),
            ("Precision/compute", metadata.get("precision_compute", "")),
            ("Audio optimizado", metadata.get("audio_optimizado", "")),
            ("Hilos CPU/FFmpeg", metadata.get("hilos_cpu_ffmpeg", "")),
            ("Troceo archivo", metadata.get("troceo_archivo", "")),
            ("Limpieza basica", metadata.get("limpieza_basica", "")),
            ("Fecha y hora", metadata.get("fecha_hora", "")),
        ]:
            row = table.add_row().cells
            row[0].text = key
            row[1].text = str(value)

        doc.add_paragraph("")
        h1 = doc.add_paragraph()
        heading = "TRANSCRIPCION EN FORMATO PREGUNTA/RESPUESTA" if "P/R" in str(metadata.get("formato", "")) else "TRANSCRIPCION COMPLETA"
        h1.add_run(heading).bold = True
        doc.add_paragraph(text)

        doc.add_paragraph("")
        h2 = doc.add_paragraph()
        h2.add_run("SEGMENTOS CON HORARIOS").bold = True
        for seg in segments:
            start = seconds_to_hhmmss(seg.get("start", 0))
            end = seconds_to_hhmmss(seg.get("end", 0))
            segment_text = str(seg.get("text", "")).strip()
            p = doc.add_paragraph()
            p.add_run(f"[{start} - {end}] ").bold = True
            p.add_run(segment_text)

        doc.save(str(path))


def main() -> int:
    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    window = MainWindow()
    window.show()
    return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())
